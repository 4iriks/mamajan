"""Calculated components use the same name/size columns as project extras."""

from copy import deepcopy
from io import BytesIO
import json
import re
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from engine.office_xlsx import _build_delivery_xlsx
from engine.project_documents import (
    _build_delivery_context,
    _delivery_hardware_groups,
    render_project_document_html,
)
from tests.test_oct29_fixes import narrow_cs


def cs_section(quantity=2):
    section = narrow_cs()
    section.system = "ЦС"
    section.name = "Секция 1"
    section.quantity = quantity
    return section


@pytest.mark.parametrize("quantities", [(2,), (1, 3)])
def test_clamp_group_preserves_summed_sizes_and_saved_places(quantities):
    project = SimpleNamespace(number="QA", production_stages=2, current_stage=1)
    sections = [cs_section(qty) for qty in quantities]
    before = deepcopy([vars(section) for section in sections])
    initial = _build_delivery_context(project, sections)
    clamps = [r for r in initial["delivery_item2_rows"] if r["article"] == "Т40Т"]
    assert len(clamps) == 4
    expected = {"3000 мм": 2, "1499 мм": 1, "1594 мм": 1, "967 мм": 1}
    assert {r["size"]: r["qty"] for r in clamps} == {
        size: qty * sum(quantities) for size, qty in expected.items()
    }
    places = {r["place_key"]: str(i + 1) for i, r in enumerate(clamps)}
    project.delivery_note_data = json.dumps({"places": places})
    context = _build_delivery_context(project, sections)
    groups = [
        g for g in context["delivery_hardware_groups"] if g[0]["article"] == "Т40Т"
    ]
    assert len(groups) == 1
    assert [r["place_key"] for r in groups[0]] == [r["place_key"] for r in clamps]
    assert [r["places"] for r in groups[0]] == ["1", "2", "3", "4"]
    assert not any(r["article"] == "Т40Т" for r in context["delivery_item2_plain_rows"])
    for is_pdf in [False, True]:
        html = render_project_document_html(
            project, sections, "delivery", is_pdf=is_pdf
        )
        assert html.count("Зажимной профиль Т40Т в сборе") == 1
        assert '<td class="extra-product" rowspan="4">' in html
        if not is_pdf:
            for key in places:
                assert html.count(f'data-delivery-place-key="{key}"') == 1
    sheet = load_workbook(BytesIO(_build_delivery_xlsx(context))).active
    heading = next(
        c for row in sheet for c in row if "Зажимной профиль" in str(c.value)
    )
    assert any(
        r.min_row == heading.row
        and r.max_row == heading.row + 3
        and r.min_col == 2
        and r.max_col == 3
        for r in sheet.merged_cells.ranges
    )
    assert [
        (
            sheet.cell(heading.row + i, 4).value,
            sheet.cell(heading.row + i, 5).value,
            sheet.cell(heading.row + i, 6).value,
        )
        for i in range(4)
    ] == [(r["size"], r["qty"], r["places"]) for r in groups[0]]
    assert [vars(section) for section in sections] == before
    project.current_stage = 2
    assert not _build_delivery_context(project, sections)["delivery_hardware_groups"]


@pytest.mark.parametrize(
    "field",
    ["article", "name", "color", "stage", "unit", "note", "image", "_view_group"],
)
def test_different_components_are_not_merged(field):
    first = dict(
        article="PROFILE",
        name="Профиль",
        color="9016",
        stage="1",
        unit="шт",
        note="",
        image="",
        _view_group="",
        size="1000 мм",
        qty=2,
        qty_text="2",
        qty_display="2 шт",
        place_key="first",
        places="1",
    )
    second = {**first, "size": "2000 мм", "place_key": "second", "places": "2"}
    plain, groups = _delivery_hardware_groups([first, second])
    assert not plain and len(groups) == 1 and len(groups[0]) == 2
    second[field] = "different"
    assert len(_delivery_hardware_groups([first, second])[1]) == 2


def test_unsized_components_keep_quantity_units_notes_and_places():
    row = dict(
        article="PAD",
        name="Подкладка",
        size="",
        unit="шт",
        qty=4,
        qty_text="4",
        note="Примечание",
        place_key="pad",
        places="2",
    )
    before = deepcopy(row)
    assert _delivery_hardware_groups([row]) == ([row], [])
    assert row == before


def test_all_delivery_components_share_one_table_and_sequential_numbers():
    project = SimpleNamespace(
        number="QA",
        production_stages=2,
        current_stage=1,
        extra_components=json.dumps([dict(name="Угол", size="1500", qty=2, unit="шт")]),
    )
    sections = [cs_section()]
    context = _build_delivery_context(project, sections)
    groups = context["delivery_component_groups"]
    first = len(context["delivery_item1_rows"]) + 1
    assert [g["number"] for g in groups] == list(range(first, first + len(groups)))
    assert context["delivery_names_count"] == first - 1 + len(groups)
    html = render_project_document_html(project, sections, "delivery")
    assert html.count('<table class="main-table component-table">') == 1
    table = html.split('<table class="main-table component-table">')[1].split(
        "</table>"
    )[0]
    assert table.count("<thead>") == 1
    assert [
        int(n)
        for n in re.findall(r'class="item-number" rowspan="\d+">(\d+)</td>', table)
    ] == [g["number"] for g in groups]
    for group in groups:
        for row in group["rows"]:
            assert table.count(f'data-delivery-place-key="{row["place_key"]}"') == 1
    sheet = load_workbook(BytesIO(_build_delivery_xlsx(context))).active
    for group in groups:
        name = group["rows"][0]["name"]
        cell = next(c for row in sheet for c in row if name in str(c.value))
        assert sheet.cell(cell.row, 1).value == group["number"]
