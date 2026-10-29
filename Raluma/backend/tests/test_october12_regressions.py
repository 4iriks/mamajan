"""Regressions limited to the October 04/08/12 amendments."""

from copy import deepcopy
import io
import re
from types import SimpleNamespace

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from engine.pdf import generate_pdf
from engine.office_docx import build_project_docx
from engine.office_xlsx import build_project_xlsx
from engine.project_documents import (
    build_project_document_context,
    render_project_document_html,
)
from engine.quote_pricing import _section_snapshot, safe_public_payload
from engine.side_projection import (
    side_projection,
    side_segments,
    snapshot_side_projection,
)
from engine.slide_calc import calculate_slide
from schemas import SectionCreate


def section(**kwargs):
    return SimpleNamespace(
        **SectionCreate(
            **{
                "name": "Проверка",
                "system": "СЛАЙД",
                "width": 3000,
                "height": 2400,
                "panels": 4,
                "quantity": 3,
                "rails": 3,
                "slide_rows": 2,
                **kwargs,
            }
        ).model_dump()
    )


@pytest.mark.parametrize("rows", [1, 2])
@pytest.mark.parametrize(
    "wall,lock,p,handle,bubble,expected",
    [
        (True, True, False, True, False, [16, 50, 14, 13]),
        (False, True, False, True, False, [50, 14, 13]),
        (True, False, True, True, False, [16, 18, 14, 13]),
        (False, False, True, True, False, [18, 14, 13]),
        (True, False, True, False, False, [16, 25]),
        (False, False, False, False, True, [10]),
        (False, False, False, False, False, []),
    ],
)
def test_visible_assemblies_do_not_change_glass(
    rows, wall, lock, p, handle, bubble, expected
):
    s = section(
        slide_rows=rows,
        profile_left_wall=wall,
        profile_left_lock_bar=lock,
        profile_left_p_bar=p,
        profile_left_handle_bar=handle,
        profile_left_bubble=bubble,
        profile_right_wall=False,
        profile_right_bubble=True,
    )
    calc = calculate_slide(s)
    before = deepcopy(calc.glass)
    segments = side_segments(s)
    assert [part["width_mm"] for part in segments["left"]] == expected
    assert [part["offset_mm"] for part in segments["left"]] == [
        sum(expected[:i]) for i in range(len(expected))
    ]
    assert side_projection(s) == {"left": sum(expected), "right": 10}
    snapshot = _section_snapshot(s, calc)
    assert snapshot["side_projection_segments"] == segments
    assert snapshot_side_projection(snapshot) == side_projection(s)
    assert calc.glass == before


@pytest.mark.parametrize(
    "stages,stage,present", [(1, 1, True), (2, 1, False), (2, 2, True)]
)
def test_rs206_delivery_uses_calculated_quantity_in_pdf_and_xlsx(
    stages, stage, present
):
    project = SimpleNamespace(
        number="QA", production_stages=stages, current_stage=stage
    )
    s = section(center_lock="Накидная защёлка RS206")
    actual = next(
        row.value for row in calculate_slide(s).hardware if row.article == "RS206"
    )
    context = build_project_document_context(project, [s], "delivery")
    found = [row for row in context["delivery_item2_rows"] if row["article"] == "RS206"]
    assert bool(found) == present
    if present:
        assert found[0]["qty"] == actual == 3
    pdf = PdfReader(
        io.BytesIO(
            generate_pdf(
                render_project_document_html(project, [s], "delivery", is_pdf=True)
            )
        )
    )
    text = "\n".join(page.extract_text() for page in pdf.pages)
    workbook = load_workbook(io.BytesIO(build_project_xlsx(project, [s], "delivery")))
    cells = "\n".join(
        str(cell.value)
        for sheet in workbook
        for row in sheet
        for cell in row
        if cell.value is not None
    )
    assert bool(re.search(r"RS206(?![0-9])", text)) == present
    assert ("RS206" in cells) == present


def quote_fixture(version=2):
    return {
        "document_layout_version": version,
        "lines": [
            {
                "name": "Тестовая услуга",
                "category": "service",
                "quantity": "1",
                "unit": "шт.",
                "document_unit_final_price": 630,
                "document_line_total": 630,
                "document_line_total_before_discount": 1000,
                "document_line_discount_amount": 370,
            }
        ],
        "payment_terms": "Оплата 70/30",
        "revision": 1,
        "totals": {
            "document_grand_total": 630,
            "document_before_discount": 1000,
            "document_discount": 370,
        },
    }


def all_docx_text(data):
    doc = Document(io.BytesIO(data))
    return "\n".join(
        [p.text for p in doc.paragraphs]
        + [
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        ]
    )


@pytest.mark.parametrize(
    "kind,version,show_discount",
    [("commercial", 2, False), ("contract_appendix", 2, True), ("commercial", 1, True)],
)
def test_discount_details_only_contract_new_quotes_and_old_snapshot_preserved(
    kind, version, show_discount
):
    project = SimpleNamespace(number="QA", customer="Проверка")
    quote = quote_fixture(version)
    before = deepcopy(quote)
    html = render_project_document_html(project, [], kind, is_pdf=True, quote=quote)
    docx = all_docx_text(build_project_docx(project, [], kind, quote=quote))
    for text in (html, docx):
        assert ("До скидки" in text) == show_discount
        assert ("Оплата 70/30" in text) == show_discount
        assert "630" in text
    assert quote == before
    assert safe_public_payload(quote)["document_layout_version"] == version


def test_sketch_extras_use_original_shared_table():
    project = SimpleNamespace(
        number="QA",
        extra_components=[
            dict(sku=sku, name=name, unit="шт.", qty=2, size=str(length))
            for sku, name, length in [
                ("EXTRA-A", "Профиль А", 1000),
                ("EXTRA-A", "Профиль А", 2000),
                ("EXTRA-B", "Профиль Б", 1500),
            ]
        ],
    )
    html = render_project_document_html(project, [], "sketch", is_pdf=True)
    assert html.count('<table class="component-table">') == 1
    assert "<th>Артикул</th><th>Наименование</th>" in html
    document = Document(io.BytesIO(build_project_docx(project, [], "sketch")))
    tables = [table for table in document.tables if table.cell(0, 0).text == "Артикул"]
    assert len(tables) == 1
    assert len(tables[0].rows) == 4
    assert tables[0].cell(1, 0).text == "EXTRA-A"
    assert tables[0].cell(2, 0).text == "EXTRA-A"
    assert tables[0].cell(3, 0).text == "EXTRA-B"


def test_delivery_repeats_group_heading_on_pages_without_mutating_sources():
    extra = dict(
        sku="QA-LONG",
        name="Длинная группа",
        color="9005",
        finishName="RAL стандарт",
        unit="м",
        deliveryStage="both",
        qty=2,
    )
    rows = [{**extra, "size": str(1000 + i * 10)} for i in range(90)]
    project = SimpleNamespace(number="QA", extra_components=rows)
    before = deepcopy(rows)
    context = build_project_document_context(project, [], "delivery")
    assert len(context["delivery_extra_groups"]) == 1
    assert all(
        row["display_unit"] == "шт." and row["pdf_qty_display"] == "2 шт."
        for row in context["delivery_project_extra_rows"]
    )
    html = render_project_document_html(project, [], "delivery", is_pdf=True)
    assert html.count('<table class="main-table">') == 1
    assert html.count('<table class="main-table extra-table">') == 4
    assert html.count("QA-LONG") == 4
    assert html.count('<td class="extra-size">') == 90
    pdf = PdfReader(io.BytesIO(generate_pdf(html)))
    pages = [page.extract_text() for page in pdf.pages]
    assert len(pages) >= 3
    assert all("QA-LONG" in page for page in pages if " мм" in page)
    workbook = load_workbook(io.BytesIO(build_project_xlsx(project, [], "delivery")))
    text = "\n".join(
        str(cell.value)
        for sheet in workbook
        for row in sheet
        for cell in row
        if cell.value is not None
    )
    assert "мм · м" not in text and "шт шт" not in text
    assert rows == before
