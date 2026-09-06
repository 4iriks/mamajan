import io
import json
import re
import zipfile
from datetime import datetime, timedelta
from decimal import Decimal
from uuid import uuid4

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from pypdf import PdfReader, PdfWriter

import models
from database import SessionLocal
from engine.quote_pricing import _section_requirements, freeze_quote


def _unique(prefix: str) -> str:
    return f"{prefix}-{uuid4().hex[:10]}"


def _login_headers(client, username: str, password: str) -> dict[str, str]:
    response = client.post(
        "/api/auth/login",
        json={"username": username, "password": password},
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _create_user(client, admin_headers, role="user", **overrides):
    password = overrides.pop("password", "secret123")
    username = overrides.pop("username", _unique(role))
    payload = {
        "username": username,
        "display_name": overrides.pop("display_name", username),
        "password": password,
        "role": role,
        "is_active": True,
        **overrides,
    }
    response = client.post("/api/users", headers=admin_headers, json=payload)
    assert response.status_code == 201, response.text
    return response.json(), password


def _new_catalog_item():
    db = SessionLocal()
    try:
        item = models.CatalogItem(
            sku=_unique("PRICE"),
            name="Тестовая позиция цены",
            group="Профили",
            system="СЛАЙД",
            unit="п.м.",
            is_active=True,
        )
        db.add(item)
        db.commit()
        db.refresh(item)
        return item.id, item.sku
    finally:
        db.close()


def _delete_catalog_item(item_id: int):
    db = SessionLocal()
    try:
        item = db.get(models.CatalogItem, item_id)
        if item is not None:
            db.delete(item)
            db.commit()
    finally:
        db.close()


def _price_payload(cost: str, effective_from: datetime, reason="Тест цены"):
    return {
        "cost": cost,
        "profile_markup_percent": "100",
        "profile_discount_percent": "25",
        "waste_markup_percent": "30",
        "construction_markup_percent": "200",
        "construction_discount_percent": "35",
        "category": "profile",
        "unit": "п.м.",
        "min_margin_percent": "10",
        "effective_from": effective_from.isoformat(),
        "reason": reason,
    }


def test_public_catalog_supports_multiple_admins(client, admin_headers):
    second_admin, _ = _create_user(client, admin_headers, role="admin")
    try:
        response = client.get("/api/catalog/hardware/options")
        assert response.status_code == 200, response.text
        assert response.json()
    finally:
        client.delete(f"/api/users/{second_admin['id']}", headers=admin_headers)


def test_price_versions_future_history_bulk_and_rollback(client, admin_headers):
    item_id, sku = _new_catalog_item()
    try:
        active = client.post(
            f"/api/pricing/catalog/{item_id}/versions",
            headers=admin_headers,
            json=_price_payload("100", datetime.utcnow() - timedelta(minutes=1)),
        )
        assert active.status_code == 201, active.text
        active_id = active.json()["id"]

        future = client.post(
            f"/api/pricing/catalog/{item_id}/versions",
            headers=admin_headers,
            json=_price_payload("200", datetime.utcnow() + timedelta(days=1)),
        )
        assert future.status_code == 201, future.text

        catalog = client.get("/api/pricing/catalog", headers=admin_headers)
        assert catalog.status_code == 200
        row = next(item for item in catalog.json()["items"] if item["id"] == item_id)
        assert row["active_price"]["cost"] == "100.00"
        assert row["next_price"]["cost"] == "200.00"
        assert row["history_count"] == 2

        bulk_data = {
            "item_ids": [item_id],
            "percent": "10",
            "effective_from": datetime.utcnow().isoformat(),
            "reason": "Индексация",
        }
        preview = client.post(
            "/api/pricing/catalog/bulk/preview",
            headers=admin_headers,
            json=bulk_data,
        )
        assert preview.status_code == 200, preview.text
        assert preview.json()["rows"][0] == {
            "item_id": item_id,
            "sku": sku,
            "name": "Тестовая позиция цены",
            "old_cost": "100.00",
            "new_cost": "110.00",
            "source_version_id": active_id,
        }
        applied = client.post(
            "/api/pricing/catalog/bulk/apply",
            headers=admin_headers,
            json=bulk_data,
        )
        assert applied.status_code == 201, applied.text
        assert applied.json()["versions"][0]["cost"] == "110.00"

        rolled_back = client.post(
            f"/api/pricing/catalog/{item_id}/rollback/{active_id}",
            headers=admin_headers,
            json={
                "effective_from": datetime.utcnow().isoformat(),
                "reason": "Возврат проверенной цены",
            },
        )
        assert rolled_back.status_code == 201, rolled_back.text
        assert rolled_back.json()["cost"] == "100.00"
        assert rolled_back.json()["rollback_of_id"] == active_id

        history = client.get(
            f"/api/pricing/catalog/{item_id}/versions",
            headers=admin_headers,
        )
        assert history.status_code == 200
        assert len(history.json()["versions"]) == 4
    finally:
        _delete_catalog_item(item_id)


def test_standalone_sale_uses_only_item_terms_variants_and_buyer_discount(
    client, admin_headers
):
    actor_id = client.get("/api/auth/me", headers=admin_headers).json()["id"]
    db = SessionLocal()
    try:
        regular = models.CatalogItem(
            sku=_unique("SALE"),
            name="Самостоятельный профиль",
            group="Профили",
            system="СЛАЙД",
            unit="п.м.",
            purchase_price=1,
            markup_percent=1,
            waste_percent=90,
            is_active=True,
        )
        colored = models.CatalogItem(
            sku=_unique("SALE-COLOR"),
            name="Цветной профиль",
            group="Профили",
            system="СЛАЙД",
            unit="п.м.",
            purchase_price=999,
            markup_percent=999,
            waste_percent=99,
            is_active=True,
        )
        colored.finish_variants.extend(
            [
                models.CatalogFinishVariant(
                    name="Анод",
                    price=200,
                    cost=200,
                    profile_markup_percent=20,
                    profile_discount_percent=10,
                    requires_paint=False,
                ),
                models.CatalogFinishVariant(
                    name="RAL 9005",
                    price=250,
                    cost=250,
                    profile_markup_percent=20,
                    profile_discount_percent=10,
                    requires_paint=True,
                ),
            ]
        )
        db.add_all([regular, colored])
        db.flush()
        db.add(
            models.CatalogPriceVersion(
                catalog_item_id=regular.id,
                cost=100,
                profile_markup_percent=20,
                profile_discount_percent=10,
                waste_markup_percent=80,
                construction_markup_percent=300,
                construction_discount_percent=75,
                category="profile",
                unit="п.м.",
                min_margin_percent=0,
                effective_from=datetime.utcnow() - timedelta(minutes=1),
                created_by=actor_id,
                reason="pytest standalone sale",
            )
        )
        db.add(
            models.CatalogPriceVersion(
                catalog_item_id=colored.id,
                cost=999,
                profile_markup_percent=20,
                profile_discount_percent=10,
                waste_markup_percent=90,
                construction_markup_percent=300,
                construction_discount_percent=75,
                category="profile",
                unit="п.м.",
                min_margin_percent=0,
                effective_from=datetime.utcnow() - timedelta(minutes=1),
                created_by=actor_id,
                reason="pytest standalone finish price",
            )
        )
        db.commit()
        db.refresh(regular)
        db.refresh(colored)
        regular_id = regular.id
        colored_id = colored.id
        ral_id = next(
            row.id for row in colored.finish_variants if row.name == "RAL 9005"
        )
    finally:
        db.close()

    try:
        percent = client.post(
            "/api/pricing/sale/quote",
            headers=admin_headers,
            json={
                "items": [{"catalog_item_id": regular_id, "quantity": 2}],
                "buyer_discount_mode": "percent",
                "buyer_discount_value": 10,
            },
        )
        assert percent.status_code == 200, percent.text
        assert percent.json()["totals"] == {
            "before_discount": "216.00",
            "discount": "21.60",
            "grand_total": "194.40",
        }

        fixed = client.post(
            "/api/pricing/sale/quote",
            headers=admin_headers,
            json={
                "items": [{"catalog_item_id": regular_id, "quantity": 2}],
                "buyer_discount_mode": "fixed",
                "buyer_discount_value": 16,
            },
        )
        assert fixed.status_code == 200
        assert fixed.json()["totals"]["grand_total"] == "200.00"

        variant = client.post(
            "/api/pricing/sale/quote",
            headers=admin_headers,
            json={
                "items": [
                    {
                        "catalog_item_id": colored_id,
                        "finish_variant_id": ral_id,
                        "quantity": 2,
                    }
                ],
                "buyer_discount_mode": "percent",
                "buyer_discount_value": 10,
            },
        )
        assert variant.status_code == 200, variant.text
        assert variant.json()["totals"]["before_discount"] == "540.00"
        assert variant.json()["totals"]["grand_total"] == "486.00"
        assert variant.json()["lines"][0]["component_details"]["finish"] == ("RAL 9005")

        missing_variant = client.post(
            "/api/pricing/sale/quote",
            headers=admin_headers,
            json={"items": [{"catalog_item_id": colored_id, "quantity": 1}]},
        )
        assert missing_variant.status_code == 400
        assert "исполнение" in missing_variant.json()["detail"].lower()

        serialized = json.dumps(percent.json(), ensure_ascii=False).lower()
        for forbidden in (
            "cost",
            "purchase",
            "margin",
            "waste",
            "production",
            "cut",
            "glue",
        ):
            assert forbidden not in serialized
    finally:
        _delete_catalog_item(regular_id)
        _delete_catalog_item(colored_id)


def test_construction_price_groups_are_admin_managed_and_publicly_safe(
    client, admin_headers
):
    admin_list = client.get("/api/pricing/price-groups", headers=admin_headers)
    assert admin_list.status_code == 200
    assert {row["code"] for row in admin_list.json()} == {"SLIDE_1", "SLIDE_2"}

    public_list = client.get("/api/catalog/construction-price-groups")
    assert public_list.status_code == 200
    assert {row["code"] for row in public_list.json()} == {"SLIDE_1", "SLIDE_2"}
    assert all("markup_percent" not in row for row in public_list.json())

    fixed = client.get("/api/catalog/system-markups", headers=admin_headers)
    assert fixed.status_code == 200
    assert {row["code"] for row in fixed.json()} == {"SLIDE_1", "SLIDE_2"}

    rejected = client.post(
        "/api/pricing/price-groups",
        headers=admin_headers,
        json={
            "code": _unique("FUTURE").upper(),
            "name": "Будущая система",
            "markup_percent": "33.5",
            "is_active": True,
        },
    )
    assert rejected.status_code == 409

    assert client.get("/api/pricing/settings", headers=admin_headers).status_code == 404


def test_unified_catalog_uses_fixed_finish_rows_and_versioned_prices(
    client, admin_headers
):
    sku = _unique("UNIFIED").upper()
    item_id = None
    payload = {
        "sku": sku,
        "name": "Профиль единого каталога",
        "group": "Профили",
        "system": "СЛАЙД",
        "systemGroups": ["SLIDE_1", "SLIDE_2"],
        "unit": "м.п.",
        "purchasePrice": 100,
        "markupPercent": 10,
        "profileDiscountPercent": 5,
        "weight": 0.5,
        "wastePercent": 7,
        "constructionMarkupPercent": 20,
        "constructionDiscountPercent": 8,
        "sectionWidthMm": 72,
        "sectionHeightMm": 53,
        "imageFile": "",
        "paintMode": "Красится",
        "colorVariants": [],
        "finishVariants": [
            {
                "code": "ANOD",
                "name": "Имя от клиента игнорируется",
                "cost": 110,
                "profileMarkupPercent": 11,
                "profileDiscountPercent": 1,
                "constructionMarkupPercent": 21,
                "constructionDiscountPercent": 2,
                "requiresPaint": True,
                "isActive": True,
            },
            {
                "code": "RAL_STANDARD",
                "name": "Другое имя",
                "cost": 120,
                "profileMarkupPercent": 12,
                "profileDiscountPercent": 3,
                "constructionMarkupPercent": 22,
                "constructionDiscountPercent": 4,
                "requiresPaint": False,
                "isActive": True,
            },
            {
                "code": "RAL_NONSTANDARD",
                "name": "Еще одно имя",
                "cost": 130,
                "profileMarkupPercent": 13,
                "profileDiscountPercent": 5,
                "constructionMarkupPercent": 23,
                "constructionDiscountPercent": 6,
                "requiresPaint": False,
                "isActive": True,
            },
        ],
        "supplier": "",
        "isActive": True,
        "note": "",
    }
    try:
        created = client.post(
            "/api/catalog/hardware", headers=admin_headers, json=payload
        )
        assert created.status_code == 201, created.text
        item = created.json()
        item_id = item["id"]
        assert item["systemGroups"] == ["SLIDE_1", "SLIDE_2"]
        assert [row["code"] for row in item["finishVariants"]] == [
            "ANOD",
            "RAL_STANDARD",
            "RAL_NONSTANDARD",
        ]
        assert [row["name"] for row in item["finishVariants"]] == [
            "Анод",
            "RAL стандарт",
            "RAL нестандарт",
        ]
        assert [row["requiresPaint"] for row in item["finishVariants"]] == [
            False,
            True,
            True,
        ]
        assert [float(row["cost"]) for row in item["finishVariants"]] == [
            110,
            120,
            130,
        ]

        db = SessionLocal()
        try:
            saved = db.get(models.CatalogItem, item_id)
            assert len(saved.finish_variants) == 3
            assert len(saved.price_versions) == 3
            assert all(row.finish_variant_id for row in saved.price_versions)
        finally:
            db.close()

        base_payload = {
            **payload,
            "paintMode": "Не красится",
            "finishVariants": [
                {
                    "code": "BASE",
                    "name": "Произвольное имя",
                    "cost": 95,
                    "profileMarkupPercent": 9,
                    "profileDiscountPercent": 1,
                    "constructionMarkupPercent": 19,
                    "constructionDiscountPercent": 2,
                    "requiresPaint": True,
                    "isActive": True,
                }
            ],
        }
        updated = client.put(
            f"/api/catalog/hardware/{item_id}",
            headers=admin_headers,
            json=base_payload,
        )
        assert updated.status_code == 200, updated.text
        assert [row["code"] for row in updated.json()["finishVariants"]] == ["BASE"]
        assert updated.json()["finishVariants"][0]["name"] == "Без окраски"
        assert updated.json()["finishVariants"][0]["requiresPaint"] is False

        standalone = client.post(
            "/api/pricing/sale/quote",
            headers=admin_headers,
            json={"items": [{"catalog_item_id": item_id, "quantity": 1}]},
        )
        assert standalone.status_code == 200, standalone.text

        db = SessionLocal()
        try:
            saved = db.get(models.CatalogItem, item_id)
            assert len(saved.finish_variants) == 4
            assert {row.code for row in saved.finish_variants if row.is_active} == {
                "BASE"
            }
            assert len(saved.price_versions) == 4
        finally:
            db.close()
    finally:
        if item_id is not None:
            _delete_catalog_item(item_id)


def test_excel_import_reads_formatted_percentages_and_applies_atomically(
    client,
    admin_headers,
):
    item_id, sku = _new_catalog_item()
    try:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Цены"
        sheet.append(
            [
                "Артикул",
                "Себестоимость",
                "Наценка на профиль",
                "Скидка на профиль",
                "Наценка на отходы",
                "Наценка на конструкции",
                "Скидка на конструкции",
                "Категория",
                "Единица",
                "Минимальная маржа",
            ]
        )
        sheet.append([sku, 123.45, 1, 0.25, 0.3, 2, 0.35, "profile", "п.м.", 0.1])
        for cell in sheet[2][2:7]:
            cell.number_format = "0%"
        sheet["J2"].number_format = "0%"
        output = io.BytesIO()
        workbook.save(output)

        preview = client.post(
            "/api/pricing/catalog/import/preview",
            headers=admin_headers,
            files={
                "file": (
                    "prices.xlsx",
                    output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert preview.status_code == 200, preview.text
        data = preview.json()
        assert data["valid"] is True
        assert data["rows"][0]["profile_markup_percent"] == "100"
        assert data["rows"][0]["profile_discount_percent"] == "25"
        assert data["rows"][0]["construction_markup_percent"] == "200"
        assert data["rows"][0]["min_margin_percent"] == "10"

        invalid_atomic = client.post(
            "/api/pricing/catalog/import/apply",
            headers=admin_headers,
            json={
                "rows": [*data["rows"], {**data["rows"][0], "sku": "NOT-FOUND"}],
                "reason": "Проверка атомарности",
            },
        )
        assert invalid_atomic.status_code == 400
        history = client.get(
            f"/api/pricing/catalog/{item_id}/versions",
            headers=admin_headers,
        ).json()
        assert history["versions"] == []

        applied = client.post(
            "/api/pricing/catalog/import/apply",
            headers=admin_headers,
            json={"rows": data["rows"], "reason": "Импорт прайс-листа"},
        )
        assert applied.status_code == 201, applied.text
        assert applied.json()["versions"][0]["cost"] == "123.45"
    finally:
        _delete_catalog_item(item_id)


def test_supplier_cost_import_updates_only_anod_or_base_and_skips_pending(
    client,
    admin_headers,
):
    paint_sku = _unique("COST-PAINT").upper()
    base_sku = _unique("COST-BASE").upper()
    new_sku = _unique("COST-NEW").upper()
    touched_skus = {paint_sku, base_sku, new_sku, "RS1006", "RS3110", "RU003"}
    seeded = client.get("/api/catalog/hardware", headers=admin_headers)
    assert seeded.status_code == 200, seeded.text
    db = SessionLocal()
    existing_version_ids: set[int] = set()
    original_items: dict[str, dict] = {}
    try:
        paint_item = models.CatalogItem(
            sku=paint_sku,
            name="Тестовый окрашиваемый профиль",
            group="Профили",
            system="СЛАЙД",
            unit="м.п.",
            purchase_price=10,
            markup_percent=35,
            waste_percent=4,
            paint_mode="Красится",
            is_active=True,
        )
        for code, name, cost, requires_paint in (
            ("ANOD", "Анод", 10, False),
            ("RAL_STANDARD", "RAL стандарт", 222.22, True),
            ("RAL_NONSTANDARD", "RAL нестандарт", 333.33, True),
        ):
            paint_item.finish_variants.append(
                models.CatalogFinishVariant(
                    code=code,
                    name=name,
                    cost=cost,
                    price=cost,
                    profile_markup_percent=35,
                    profile_discount_percent=0,
                    construction_markup_percent=5,
                    construction_discount_percent=0,
                    requires_paint=requires_paint,
                    is_active=True,
                )
            )
        base_item = models.CatalogItem(
            sku=base_sku,
            name="Тестовая защёлка",
            group="Защёлки",
            system="СЛАЙД",
            unit="шт",
            purchase_price=20,
            markup_percent=40,
            paint_mode="Не красится",
            is_active=True,
        )
        base_item.finish_variants.append(
            models.CatalogFinishVariant(
                code="BASE",
                name="Без окраски",
                cost=20,
                price=20,
                profile_markup_percent=40,
                profile_discount_percent=0,
                construction_markup_percent=0,
                construction_discount_percent=0,
                requires_paint=False,
                is_active=True,
            )
        )
        db.add_all([paint_item, base_item])
        db.commit()

        for sku in touched_skus - {new_sku}:
            item = db.query(models.CatalogItem).filter_by(sku=sku).first()
            assert item is not None
            original_items[sku] = {
                "id": item.id,
                "purchase_price": item.purchase_price,
                "variants": {
                    variant.id: (variant.cost, variant.price)
                    for variant in item.finish_variants
                },
            }
            existing_version_ids.update(version.id for version in item.price_versions)
    finally:
        db.close()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Стоим Слайд 26"
    sheet["B5"] = "04.09.2026г."
    sheet.merge_cells("A8:E8")
    sheet["A8"] = "Наименование/Артикул профиля Слайд"
    sheet["F8"] = "Себестоимость"
    supplier_rows = (
        (paint_sku, "Тестовый окрашиваемый профиль", "м.п.", 101.25),
        (base_sku, "Тестовая защёлка", "шт", 45.50),
        (new_sku, "Новый направляющий профиль", "м.п.", 77.70),
        ("RS1006", "Прозрачный межстворочный уплотнитель", "м.п.", 390),
        ("RS3110", "h-уплотнитель центрального стыка", "м.п.", 710),
        ("RU003", "Ролик 2-колесный", "шт", 250),
    )
    for row_number, values in enumerate(supplier_rows, start=10):
        for column, value in zip(("C", "D", "E", "F"), values):
            sheet[f"{column}{row_number}"] = value
    sheet["F15"].fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    output = io.BytesIO()
    workbook.save(output)

    try:
        preview_response = client.post(
            "/api/pricing/catalog/import/preview",
            headers=admin_headers,
            files={
                "file": (
                    "supplier-cost.xlsx",
                    output.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert preview_response.status_code == 200, preview_response.text
        preview = preview_response.json()
        assert preview["valid"] is True
        assert preview["can_apply"] is True
        rows = {row["sku"]: row for row in preview["rows"]}
        assert rows[paint_sku]["finish_code"] == "ANOD"
        assert rows[base_sku]["finish_code"] == "BASE"
        assert rows[new_sku]["action"] == "create"
        assert rows[new_sku]["finish_code"] == "ANOD"
        assert rows["RS1006"]["cost"] == "125.81"
        assert rows["RS3110"]["cost"] == "236.67"
        assert rows["RU003"]["status"] == "pending"
        assert rows["RU003"]["action"] == "skip"

        applied = client.post(
            "/api/pricing/catalog/import/apply",
            headers=admin_headers,
            json={"rows": preview["rows"], "reason": "Тест сверки себестоимости"},
        )
        assert applied.status_code == 201, applied.text
        result = applied.json()
        assert any(row["sku"] == new_sku for row in result["created_items"])
        assert any(row["sku"] == "RU003" for row in result["skipped"])

        db = SessionLocal()
        try:
            refreshed_paint = db.query(models.CatalogItem).filter_by(sku=paint_sku).one()
            variant_costs = {
                row.code: Decimal(str(row.cost)) for row in refreshed_paint.finish_variants
            }
            assert variant_costs["ANOD"] == Decimal("101.25")
            assert variant_costs["RAL_STANDARD"] == Decimal("222.22")
            assert variant_costs["RAL_NONSTANDARD"] == Decimal("333.33")
            refreshed_base = db.query(models.CatalogItem).filter_by(sku=base_sku).one()
            assert Decimal(str(refreshed_base.finish_variants[0].cost)) == Decimal("45.50")
            created = db.query(models.CatalogItem).filter_by(sku=new_sku).one()
            created_variants = {row.code: Decimal(str(row.cost)) for row in created.finish_variants}
            assert created_variants == {
                "ANOD": Decimal("77.70"),
                "RAL_STANDARD": Decimal("0.00"),
                "RAL_NONSTANDARD": Decimal("0.00"),
            }
            assert Decimal(
                str(
                    db.query(models.CatalogItem)
                    .filter_by(sku="RS1006")
                    .one()
                    .finish_variants[0]
                    .cost
                )
            ) == Decimal("125.81")
            assert Decimal(
                str(
                    db.query(models.CatalogItem)
                    .filter_by(sku="RS3110")
                    .one()
                    .finish_variants[0]
                    .cost
                )
            ) == Decimal("236.67")
            pending = db.query(models.CatalogItem).filter_by(sku="RU003").one()
            assert Decimal(str(pending.finish_variants[0].cost)) == Decimal(
                str(original_items["RU003"]["variants"][pending.finish_variants[0].id][0])
            )
            anod = next(row for row in refreshed_paint.finish_variants if row.code == "ANOD")
            assert any(
                version.finish_variant_id == anod.id
                for version in refreshed_paint.price_versions
            )
        finally:
            db.close()

        repeated = client.post(
            "/api/pricing/catalog/import/apply",
            headers=admin_headers,
            json={"rows": preview["rows"], "reason": "Тест сверки себестоимости"},
        )
        assert repeated.status_code == 201, repeated.text
        assert set(repeated.json()["unchanged"]) >= {
            paint_sku,
            base_sku,
            new_sku,
            "RS1006",
            "RS3110",
        }
        assert repeated.json()["versions"] == []
    finally:
        db = SessionLocal()
        try:
            touched_ids = [data["id"] for data in original_items.values()]
            created_item = db.query(models.CatalogItem).filter_by(sku=new_sku).first()
            if created_item is not None:
                touched_ids.append(created_item.id)
            if touched_ids:
                db.query(models.CatalogPriceVersion).filter(
                    models.CatalogPriceVersion.catalog_item_id.in_(touched_ids),
                    ~models.CatalogPriceVersion.id.in_(existing_version_ids or {-1}),
                ).delete(synchronize_session=False)
            if created_item is not None:
                for variant in list(created_item.finish_variants):
                    db.delete(variant)
                db.delete(created_item)
            for sku, data in original_items.items():
                item = db.get(models.CatalogItem, data["id"])
                if item is None:
                    continue
                item.purchase_price = data["purchase_price"]
                for variant in item.finish_variants:
                    if variant.id in data["variants"]:
                        variant.cost, variant.price = data["variants"][variant.id]
            for sku in (paint_sku, base_sku):
                item = db.query(models.CatalogItem).filter_by(sku=sku).first()
                if item is not None:
                    for variant in list(item.finish_variants):
                        db.delete(variant)
                    db.delete(item)
            db.commit()
        finally:
            db.close()


def test_pricing_permission_and_dealer_quote_access(client, admin_headers):
    manager, manager_password = _create_user(client, admin_headers, role="user")
    price_manager, price_password = _create_user(
        client,
        admin_headers,
        role="user",
        can_manage_prices=True,
    )
    dealer, dealer_password = _create_user(client, admin_headers, role="dealer")
    dealer_headers = _login_headers(client, dealer["username"], dealer_password)
    manager_headers = _login_headers(client, manager["username"], manager_password)
    price_headers = _login_headers(client, price_manager["username"], price_password)
    project_id = None
    manager_project_id = None
    price_manager_project_id = None
    try:
        dealer_me = client.get("/api/auth/me", headers=dealer_headers)
        assert dealer_me.status_code == 200
        assert "dealer_discount_percent" not in dealer_me.json()
        assert (
            client.get("/api/pricing/catalog", headers=manager_headers).status_code
            == 403
        )
        assert (
            client.get("/api/pricing/catalog", headers=dealer_headers).status_code
            == 403
        )
        assert (
            client.get("/api/pricing/catalog", headers=price_headers).status_code == 403
        )
        for headers in (manager_headers, dealer_headers, price_headers):
            assert (
                client.get("/api/catalog/hardware", headers=headers).status_code == 403
            )
            assert (
                client.get("/api/catalog/system-markups", headers=headers).status_code
                == 403
            )
        terms = client.put(
            f"/api/pricing/dealers/{dealer['id']}",
            headers=admin_headers,
            json={
                "dealer_markup_percent": "20",
                "profile_discount_percent": "5",
                "construction_discount_percent": "10",
                "component_discount_percent": "7",
                "service_discount_percent": "3",
            },
        )
        assert terms.status_code == 200, terms.text
        assert terms.json()["construction_discount_percent"] == "10"
        assert (
            client.get(
                f"/api/pricing/dealers/{dealer['id']}", headers=dealer_headers
            ).status_code
            == 403
        )
        manager_project = client.post(
            "/api/projects",
            headers=manager_headers,
            json={"number": _unique("USER-Q"), "customer": "Клиент"},
        )
        assert manager_project.status_code == 201
        manager_project_id = manager_project.json()["id"]
        public_only = client.get(
            f"/api/projects/{manager_project_id}/quote",
            headers=manager_headers,
        )
        assert public_only.status_code == 200
        _assert_no_internal_pricing(public_only.json())
        manager_editor = client.get(
            f"/api/pricing/projects/{manager_project_id}",
            headers=manager_headers,
        )
        assert manager_editor.status_code == 200
        assert manager_editor.json()["calculation"] == {}
        assert manager_editor.json()["missing_prices"] == []
        assert manager_editor.json()["config"]["overrides"] == []
        config = {
            "validity_days": 14,
            "manufacturing_term": "",
            "payment_terms": "",
            "services": [],
        }
        assert (
            client.put(
                f"/api/projects/{manager_project_id}/quote/config",
                headers=manager_headers,
                json=config,
            ).status_code
            == 200
        )
        assert (
            client.post(
                f"/api/projects/{manager_project_id}/quote/refresh",
                headers=manager_headers,
            ).status_code
            == 200
        )
        price_manager_project = client.post(
            "/api/projects",
            headers=price_headers,
            json={"number": _unique("PRICE-Q"), "customer": "Клиент"},
        )
        assert price_manager_project.status_code == 201
        price_manager_project_id = price_manager_project.json()["id"]
        price_manager_editor = client.get(
            f"/api/pricing/projects/{price_manager_project_id}",
            headers=price_headers,
        )
        assert price_manager_editor.status_code == 200
        assert price_manager_editor.json()["calculation"] != {}
        assert price_manager_editor.json()["missing_prices"] == []
        price_token = price_headers["Authorization"].split(" ", 1)[1]
        manager_token = manager_headers["Authorization"].split(" ", 1)[1]
        assert (
            client.get(
                f"/api/projects/{price_manager_project_id}/documents/cost_report/preview",
                params={"token": price_token},
            ).status_code
            == 200
        )
        assert (
            client.get(
                f"/api/projects/{price_manager_project_id}/documents/cost_report/preview",
                params={"token": manager_token},
            ).status_code
            == 403
        )
        one_time_price = client.put(
            f"/api/projects/{price_manager_project_id}/quote/overrides",
            headers=price_headers,
            json={
                "overrides": [
                    {
                        "sku": "SPECIAL-PRICE",
                        "cost": "100",
                        "comment": "Согласовано для конкретного КП",
                    }
                ]
            },
        )
        assert one_time_price.status_code == 200
        margin_override = client.put(
            f"/api/projects/{price_manager_project_id}/quote/overrides",
            headers=price_headers,
            json={
                "overrides": [],
                "margin_override_comment": "Попытка разрешить исключение",
            },
        )
        assert margin_override.status_code == 403

        project = client.post(
            "/api/projects",
            headers=dealer_headers,
            json={"number": _unique("DEALER-Q"), "customer": "Дилер"},
        )
        assert project.status_code == 201, project.text
        project_id = project.json()["id"]
        assert client.get(f"/api/projects/{project_id}/quote").status_code == 403
        assert (
            client.get(
                f"/api/pricing/projects/{project_id}", headers=dealer_headers
            ).status_code
            == 403
        )
        assert (
            client.put(
                f"/api/projects/{project_id}/quote/config",
                headers=dealer_headers,
                json=config,
            ).status_code
            == 200
        )
    finally:
        if manager_project_id is not None:
            client.delete(
                f"/api/projects/{manager_project_id}", headers=manager_headers
            )
        if project_id is not None:
            client.delete(f"/api/projects/{project_id}", headers=dealer_headers)
        if price_manager_project_id is not None:
            client.delete(
                f"/api/projects/{price_manager_project_id}", headers=price_headers
            )
        db = SessionLocal()
        try:
            db.query(models.DealerPricingTerms).filter_by(user_id=dealer["id"]).delete()
            db.commit()
        finally:
            db.close()
        for user in (manager, price_manager, dealer):
            client.delete(f"/api/users/{user['id']}", headers=admin_headers)


def test_cost_report_uses_fixed_internal_revision_and_exports_pdf_xlsx(
    client,
    admin_headers,
    project,
    section,
    monkeypatch,
):
    admin_id = client.get("/api/auth/me", headers=admin_headers).json()["id"]
    public = {
        "project": {
            "id": project["id"],
            "number": project["number"],
            "invoice_number": project.get("invoice_number"),
            "order_number": project.get("order_number"),
            "customer": project["customer"],
        },
        "revision": 3,
        "status": "fixed",
        "fixed_at": "2026-09-04T12:00:00",
        "lines": [
            {"id": f"section-{section['id']}", "line_total": "500.00"},
            {"id": "project-extra-1", "line_total": "120.00"},
            {"id": "service-1", "line_total": "200.00"},
        ],
        "totals": {"grand_total": "820.00"},
        "export_allowed": True,
        "warnings": [],
    }
    internal = {
        "public": public,
        "sections": [
            {
                "section_id": section["id"],
                "name": "Секция отчёта",
                "internal_total": "400.00",
                "final_price": "500.00",
                "issues": [],
                "bom": [
                    {
                        "sku": "REPORT-A",
                        "name": "Профиль A",
                        "finish": "Анод",
                        "quantity": "2",
                        "unit": "м.п.",
                        "cost": "50.00",
                        "base_cost_total": "100.00",
                        "internal_total": "150.00",
                    },
                    {
                        "sku": "REPORT-B",
                        "name": "Фурнитура B",
                        "finish": "Без окраски",
                        "quantity": "1",
                        "unit": "шт",
                        "cost": "200.00",
                        "base_cost_total": "200.00",
                        "internal_total": "250.00",
                    },
                ],
            }
        ],
        "project_extras": [
            {
                "line_id": "project-extra-1",
                "index": 1,
                "sku": "REPORT-X",
                "name": "Доп. профиль",
                "quantity": "2",
                "unit": "шт",
                "finish": "",
                "cost": "50.00",
                "base_cost_total": "100.00",
                "internal_total": "110.00",
                "final_price": "120.00",
            }
        ],
        "services": [
            {
                "id": "service-1",
                "name": "Монтаж",
                "quantity": "1",
                "unit": "услуга",
                "base_cost": "100.00",
                "internal_total": "100.00",
                "final_price": "200.00",
            }
        ],
        "issues": [],
        "blocking_issues": [],
        "missing_prices": [],
    }
    db = SessionLocal()
    try:
        state = db.query(models.ProjectQuoteState).filter_by(project_id=project["id"]).first()
        if state is None:
            state = models.ProjectQuoteState(project_id=project["id"])
            db.add(state)
        state.revision = 3
        state.status = "fixed"
        state.public_payload = json.dumps(public, ensure_ascii=False)
        state.internal_payload = json.dumps(internal, ensure_ascii=False)
        state.source_signature = "fixed-test-snapshot"
        state.fixed_at = datetime(2026, 9, 4, 12, 0, 0)
        state.fixed_by = admin_id
        db.commit()
    finally:
        db.close()

    token = admin_headers["Authorization"].split(" ", 1)[1]
    preview = client.get(
        f"/api/projects/{project['id']}/documents/cost_report/preview",
        params={"token": token},
    )
    assert preview.status_code == 200, preview.text
    assert "Себестоимость и наценка" in preview.text
    assert "500,00 ₽" in preview.text
    assert "610,00 ₽" in preview.text
    assert "820,00 ₽" in preview.text
    assert "320,00 ₽" in preview.text
    assert "64,00 %" in preview.text
    assert "39,02 %" in preview.text

    xlsx = client.get(
        f"/api/projects/{project['id']}/documents/cost_report/xlsx",
        headers=admin_headers,
    )
    assert xlsx.status_code == 200, xlsx.text
    assert xlsx.content.startswith(b"PK")
    with zipfile.ZipFile(io.BytesIO(xlsx.content)) as archive:
        workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
        assert all(
            name in workbook_xml
            for name in ("Сводка", "Комплектующие", "Услуги", "Незаполненные")
        )
        formulas = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
        assert "IFERROR" in formulas
        assert "#REF!" not in formulas
        assert "#DIV/0!" not in formulas

    def fake_pdf(_: str) -> bytes:
        output = io.BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=842, height=595)
        writer.write(output)
        return output.getvalue()

    monkeypatch.setattr("api.documents.generate_pdf", fake_pdf)
    pdf = client.get(
        f"/api/projects/{project['id']}/documents/cost_report/pdf",
        headers=admin_headers,
    )
    assert pdf.status_code == 200, pdf.text
    reader = PdfReader(io.BytesIO(pdf.content))
    assert len(reader.pages) == 1

    local = client.post(
        "/api/projects/local/documents/cost_report/pdf",
        json={"project": {}, "sections": []},
    )
    assert local.status_code == 403
    assert (
        client.get(
            f"/api/projects/{project['id']}/documents/cost_report/docx",
            headers=admin_headers,
        ).status_code
        == 400
    )


def _seed_quote_prices(project_id: int, actor_id: int):
    db = SessionLocal()
    created_versions = []
    version_by_sku = {}
    created_items = []
    try:
        project = db.get(models.Project, project_id)
        requirements = []
        for section in project.sections:
            _, section_requirements = _section_requirements(section)
            requirements.extend(section_requirements)
        by_sku = {}
        for required in requirements:
            by_sku.setdefault(required["sku"], required)
        active_priced_skus = {
            row.sku
            for row in db.query(models.CatalogItem)
            .join(models.CatalogPriceVersion)
            .filter(models.CatalogPriceVersion.effective_from <= datetime.utcnow())
            .all()
        }
        skipped_sku = next(sku for sku in by_sku if sku not in active_priced_skus)
        margin_sku = next(sku for sku in by_sku if sku != skipped_sku)
        for sku, required in by_sku.items():
            if sku == skipped_sku:
                continue
            item = db.query(models.CatalogItem).filter_by(sku=sku).first()
            if item is None:
                item = models.CatalogItem(
                    sku=sku,
                    name=required["name"],
                    group="Тест расчёта",
                    system="СЛАЙД",
                    unit=required["unit"],
                    is_active=True,
                )
                db.add(item)
                db.flush()
                created_items.append(item.id)
            version = models.CatalogPriceVersion(
                catalog_item_id=item.id,
                cost=Decimal("10.00"),
                profile_markup_percent=Decimal("0"),
                profile_discount_percent=Decimal("0"),
                waste_markup_percent=Decimal("30"),
                construction_markup_percent=Decimal("0"),
                construction_discount_percent=Decimal("0"),
                category=required["category"],
                unit=required["unit"],
                min_margin_percent=(
                    Decimal("1000") if sku == margin_sku else Decimal("0")
                ),
                effective_from=datetime.utcnow() - timedelta(minutes=1),
                created_at=datetime.utcnow(),
                created_by=actor_id,
                reason="Интеграционный тест КП",
            )
            db.add(version)
            db.flush()
            created_versions.append(version.id)
            version_by_sku[sku] = version.id
        db.commit()
        return {
            "skipped_sku": skipped_sku,
            "margin_sku": margin_sku,
            "versions": created_versions,
            "version_by_sku": version_by_sku,
            "items": created_items,
            "requirements": by_sku,
        }
    finally:
        db.close()


def _cleanup_quote_prices(seed):
    db = SessionLocal()
    try:
        if seed["versions"]:
            db.query(models.CatalogPriceVersion).filter(
                models.CatalogPriceVersion.id.in_(seed["versions"])
            ).delete(synchronize_session=False)
        if seed["items"]:
            db.query(models.CatalogItem).filter(
                models.CatalogItem.id.in_(seed["items"])
            ).delete(synchronize_session=False)
        db.commit()
    finally:
        db.close()


def _assert_no_internal_pricing(payload):
    serialized = json.dumps(payload, ensure_ascii=False).lower()
    for forbidden in (
        "base_cost",
        "internal_total",
        "dealer_markup",
        "minimum_total",
        "min_margin",
        "price_version",
        '"bom"',
        "override_comment",
        "margin_approval",
        "context_signature",
        "approved_by",
    ):
        assert forbidden not in serialized


def test_quote_category_percent_and_order_ruble_discounts(
    client, admin_headers, project, section
):
    admin_id = client.get("/api/auth/me", headers=admin_headers).json()["id"]
    seed = _seed_quote_prices(project["id"], admin_id)
    try:
        completed = client.put(
            f"/api/projects/{project['id']}/quote/overrides",
            headers=admin_headers,
            json={
                "overrides": [
                    {
                        "sku": seed["skipped_sku"],
                        "cost": "10",
                        "comment": "Полный тестовый расчёт",
                    }
                ],
                "margin_override_comment": "",
            },
        )
        assert completed.status_code == 200, completed.text
        baseline = Decimal(completed.json()["totals"]["subtotal"])

        discounted = client.put(
            f"/api/projects/{project['id']}/quote/config",
            headers=admin_headers,
            json={
                "validity_days": 14,
                "manufacturing_term": "",
                "payment_terms": "",
                "services": [],
                "discounts": [
                    {
                        "id": "construction-percent",
                        "name": "Скидка на конструкции",
                        "scope": "construction",
                        "mode": "percent",
                        "value": "10",
                    },
                    {
                        "id": "order-rubles",
                        "name": "Скидка на заказ",
                        "scope": "order",
                        "mode": "fixed",
                        "value": "25",
                    },
                ],
            },
        )
        assert discounted.status_code == 200, discounted.text
        payload = discounted.json()
        expected = (baseline * Decimal("0.90") - Decimal("25")).quantize(
            Decimal("0.01")
        )
        assert Decimal(payload["totals"]["subtotal"]) == expected
        assert Decimal(payload["totals"]["discount"]) == baseline - expected
        assert payload["discounts"] == [
            {
                "id": "construction-percent",
                "name": "Скидка на конструкции",
                "scope": "construction",
                "mode": "percent",
                "value": "10",
            },
            {
                "id": "order-rubles",
                "name": "Скидка на заказ",
                "scope": "order",
                "mode": "fixed",
                "value": "25",
            },
        ]
        _assert_no_internal_pricing(payload)
    finally:
        _cleanup_quote_prices(seed)


def test_system_markup_bulk_updates_finish_prices_and_preserves_fixed_groups(
    client, admin_headers
):
    seeded = client.get("/api/catalog/hardware", headers=admin_headers)
    assert seeded.status_code == 200, seeded.text
    db = SessionLocal()
    try:
        assert (
            db.query(models.CatalogPriceVersion)
            .filter(models.CatalogPriceVersion.finish_variant_id.is_not(None))
            .count()
            > 0
        )
        version_ids_before = {
            row[0] for row in db.query(models.CatalogPriceVersion.id).all()
        }
        variants = (
            db.query(models.CatalogFinishVariant)
            .join(models.CatalogItem)
            .filter(models.CatalogItem.is_active == True)  # noqa: E712
            .all()
        )
        original_values = {
            row.id: Decimal(row.construction_markup_percent) for row in variants
        }
        target_markup = max(original_values.values(), default=Decimal("0")) + Decimal(
            "1.25"
        )
        group = db.query(models.ConstructionPriceGroup).filter_by(code="SLIDE_1").one()
        original_group_markup = Decimal(group.markup_percent)
    finally:
        db.close()

    try:
        response = client.put(
            "/api/catalog/system-markups/SLIDE_1",
            headers=admin_headers,
            json={"constructionMarkupPercent": str(target_markup)},
        )
        assert response.status_code == 200, response.text
        assert (
            Decimal(str(response.json()["constructionMarkupPercent"])) == target_markup
        )

        db = SessionLocal()
        try:
            changed_versions = (
                db.query(models.CatalogPriceVersion)
                .filter(~models.CatalogPriceVersion.id.in_(version_ids_before))
                .all()
            )
            assert changed_versions
            assert all(row.finish_variant_id is not None for row in changed_versions)
            assert all(
                Decimal(row.construction_markup_percent) == target_markup
                for row in changed_versions
            )
        finally:
            db.close()
    finally:
        db = SessionLocal()
        try:
            for variant_id, value in original_values.items():
                variant = db.get(models.CatalogFinishVariant, variant_id)
                if variant is not None:
                    variant.construction_markup_percent = value
            group = (
                db.query(models.ConstructionPriceGroup).filter_by(code="SLIDE_1").one()
            )
            group.markup_percent = original_group_markup
            db.query(models.CatalogPriceVersion).filter(
                ~models.CatalogPriceVersion.id.in_(version_ids_before)
            ).delete(synchronize_session=False)
            db.commit()
        finally:
            db.close()


def test_slide_without_glass_removes_glass_price_and_weight_but_keeps_geometry(
    client, admin_headers, project, section
):
    admin_id = client.get("/api/auth/me", headers=admin_headers).json()["id"]
    seed = _seed_quote_prices(project["id"], admin_id)
    try:
        with_glass = client.put(
            f"/api/projects/{project['id']}/quote/overrides",
            headers=admin_headers,
            json={
                "overrides": [
                    {
                        "sku": seed["skipped_sku"],
                        "cost": "10",
                        "comment": "Полный тестовый расчёт",
                    }
                ]
            },
        )
        assert with_glass.status_code == 200, with_glass.text
        with_payload = with_glass.json()
        with_line = next(
            row for row in with_payload["lines"] if row.get("section_details")
        )
        assert Decimal(with_line["section_details"]["glass_weight_kg"]) > 0

        section_payload = {
            key: value
            for key, value in section.items()
            if key not in {"id", "project_id"}
        }
        section_payload["glass_supplied"] = False
        updated = client.put(
            f"/api/projects/{project['id']}/sections/{section['id']}",
            headers=admin_headers,
            json=section_payload,
        )
        assert updated.status_code == 200, updated.text

        without_glass = client.get(
            f"/api/projects/{project['id']}/quote", headers=admin_headers
        )
        assert without_glass.status_code == 200
        without_payload = without_glass.json()
        without_line = next(
            row for row in without_payload["lines"] if row.get("section_details")
        )
        details = without_line["section_details"]
        assert details["glass_supplied"] is False
        assert details["glass_type"] == "Без стекла"
        assert Decimal(details["glass_weight_kg"]) == 0
        assert (
            details["panel_geometry"] == with_line["section_details"]["panel_geometry"]
        )
        assert Decimal(without_payload["totals"]["subtotal"]) < Decimal(
            with_payload["totals"]["subtotal"]
        )

        token = admin_headers["Authorization"].replace("Bearer ", "")
        preview = client.get(
            f"/api/projects/{project['id']}/documents/commercial/preview",
            params={"token": token},
        )
        assert preview.status_code == 200
        assert "Без стекла" in preview.text
        assert str(round(float(details["panel_geometry"][0]["width_mm"]))) in (
            preview.text
        )
        frame = re.search(
            r'data-quote-room-frame="[^"]+"[^>]*width="([0-9.]+)" height="([0-9.]+)"',
            preview.text,
        )
        assert frame is not None
        assert float(frame.group(1)) / float(frame.group(2)) == pytest.approx(
            2000 / 2400,
            abs=0.01,
        )
    finally:
        _cleanup_quote_prices(seed)


def test_quote_missing_override_margin_snapshot_and_safe_exports(
    client,
    admin_headers,
    project,
    section,
):
    admin_id = client.get("/api/auth/me", headers=admin_headers).json()["id"]
    initial = client.get(f"/api/projects/{project['id']}/quote", headers=admin_headers)
    assert initial.status_code == 200
    assert initial.json()["export_allowed"] is False
    assert initial.json()["missing_price_count"] > 0
    assert "missing_prices" not in initial.json()
    premature_approval = client.put(
        f"/api/projects/{project['id']}/quote/overrides",
        headers=admin_headers,
        json={
            "overrides": [],
            "margin_override_comment": "Согласование до появления нарушения",
        },
    )
    assert premature_approval.status_code == 400
    blocked_pdf = client.get(
        f"/api/projects/{project['id']}/documents/commercial/pdf",
        headers=admin_headers,
    )
    assert blocked_pdf.status_code == 409
    blocked_quote = blocked_pdf.json()["detail"]["quote"]
    assert "missing_prices" not in blocked_quote
    _assert_no_internal_pricing(blocked_quote)

    seed = _seed_quote_prices(project["id"], admin_id)
    try:
        missing = client.get(
            f"/api/projects/{project['id']}/quote", headers=admin_headers
        ).json()
        assert missing["missing_price_count"] == 1
        internal_missing = client.get(
            f"/api/pricing/projects/{project['id']}", headers=admin_headers
        ).json()
        assert {row["sku"] for row in internal_missing["missing_prices"]} == {
            seed["skipped_sku"]
        }
        token = admin_headers["Authorization"].replace("Bearer ", "")
        incomplete_preview = client.get(
            f"/api/projects/{project['id']}/documents/commercial/preview",
            params={"token": token},
        )
        assert incomplete_preview.status_code == 200
        assert seed["skipped_sku"] not in incomplete_preview.text

        invalid_override = client.put(
            f"/api/projects/{project['id']}/quote/overrides",
            headers=admin_headers,
            json={
                "overrides": [
                    {"sku": seed["skipped_sku"], "cost": "12", "comment": ""}
                ],
                "margin_override_comment": "",
            },
        )
        assert invalid_override.status_code == 422

        overridden = client.put(
            f"/api/projects/{project['id']}/quote/overrides",
            headers=admin_headers,
            json={
                "overrides": [
                    {
                        "sku": seed["skipped_sku"],
                        "cost": "12",
                        "comment": "Разовая согласованная цена",
                    }
                ],
                "margin_override_comment": "",
            },
        )
        assert overridden.status_code == 200, overridden.text
        assert overridden.json()["missing_price_count"] == 0
        assert overridden.json()["export_allowed"] is False
        assert overridden.json()["warnings"] == ["Расчёт требует проверки менеджером."]
        assert (
            "минималь" not in json.dumps(overridden.json(), ensure_ascii=False).lower()
        )
        internal_margin = client.get(
            f"/api/pricing/projects/{project['id']}", headers=admin_headers
        ).json()
        assert any(
            "минималь" in warning.lower()
            for warning in internal_margin["pending_warnings"]
        )

        unsupported_unit = client.put(
            f"/api/projects/{project['id']}/quote/config",
            headers=admin_headers,
            json={
                "validity_days": 14,
                "manufacturing_term": "",
                "payment_terms": "",
                "services": [
                    {
                        "id": "bad-unit",
                        "name": "Доставка",
                        "quantity": "1",
                        "unit": "рейс",
                        "base_cost": "100",
                    }
                ],
            },
        )
        assert unsupported_unit.status_code == 400

        configured = client.put(
            f"/api/projects/{project['id']}/quote/config",
            headers=admin_headers,
            json={
                "validity_days": 21,
                "manufacturing_term": "20 рабочих дней",
                "payment_terms": "70% аванс, 30% перед отгрузкой",
                "services": [
                    {
                        "id": "delivery",
                        "name": "Доставка",
                        "quantity": "2",
                        "unit": "шт.",
                        "base_cost": "100",
                    }
                ],
            },
        )
        assert configured.status_code == 200, configured.text

        allowed = client.put(
            f"/api/projects/{project['id']}/quote/overrides",
            headers=admin_headers,
            json={
                "overrides": [
                    {
                        "sku": seed["skipped_sku"],
                        "cost": "12",
                        "comment": "Разовая согласованная цена",
                    }
                ],
                "margin_override_comment": "Разрешено руководителем для тендера",
            },
        )
        assert allowed.status_code == 200, allowed.text
        quote = allowed.json()
        assert quote["export_allowed"] is True
        assert "vat" not in quote
        assert "vat" not in quote["totals"]
        assert quote["totals"]["grand_total"] == quote["totals"]["subtotal"]
        assert quote["totals"]["document_grand_total"] == sum(
            row["document_line_total"] for row in quote["lines"]
        )
        construction_line = next(
            row for row in quote["lines"] if row.get("section_details")
        )
        assert construction_line["section_details"]["glass_type"] == (
            "10ММ ЗАКАЛЕННОЕ ПРОЗРАЧНОЕ"
        )
        assert len(construction_line["section_details"]["panel_geometry"]) == 3
        assert construction_line["breakdown"][-1]["name"] == "Стекло"
        assert (
            sum(row["line_total"] for row in construction_line["breakdown"])
            == construction_line["document_line_total"]
        )
        _assert_no_internal_pricing(quote)
        assert "margin_approval" not in quote
        approved_internal = client.get(
            f"/api/pricing/projects/{project['id']}", headers=admin_headers
        ).json()
        approval = approved_internal["margin_approval"]
        assert approval["required"] is True
        assert approval["valid"] is True
        assert approval["target_revision"] == 1
        assert approval["approved_revision"] == 1
        assert approval["approved_by"] == admin_id
        assert approval["approved_at"]

        for _ in range(2):
            draft_refresh = client.post(
                f"/api/projects/{project['id']}/quote/refresh",
                headers=admin_headers,
            )
            assert draft_refresh.status_code == 200, draft_refresh.text
            assert draft_refresh.json()["revision"] == 1
            assert draft_refresh.json()["status"] == "draft"

        preview = client.get(
            f"/api/projects/{project['id']}/documents/commercial/preview",
            params={"token": token},
        )
        assert preview.status_code == 200, preview.text
        assert "Себестоимость" not in preview.text
        assert "Разовая согласованная цена" not in preview.text
        assert "Разрешено руководителем для тендера" not in preview.text
        assert "margin_approval" not in preview.text
        assert "70% аванс" in preview.text
        assert "Общие и центральные" in preview.text
        assert "Вид на изделие со стороны помещения" in preview.text
        assert "Вид сверху" not in preview.text
        assert "/api/catalog/profile-assets/" not in preview.text

        word_draft = client.get(
            f"/api/projects/{project['id']}/documents/commercial/docx",
            headers=admin_headers,
        )
        assert word_draft.status_code == 200, word_draft.text
        assert (
            client.get(
                f"/api/projects/{project['id']}/quote", headers=admin_headers
            ).json()["status"]
            == "draft"
        )
        with zipfile.ZipFile(io.BytesIO(word_draft.content)) as archive:
            word_xml = archive.read("word/document.xml").decode("utf-8")
            quote_images = [
                name for name in archive.namelist() if name.startswith("word/media/")
            ]
        assert "Себестоимость" not in word_xml
        assert "Разовая согласованная цена" not in word_xml
        assert "Разрешено руководителем для тендера" not in word_xml
        assert "margin_approval" not in word_xml
        assert "Стекло" in word_xml
        assert "ТЕХНИЧЕСКАЯ КОМПЛЕКТАЦИЯ" in word_xml
        assert len(quote_images) == 1

        pdf = client.get(
            f"/api/projects/{project['id']}/documents/commercial/pdf",
            headers=admin_headers,
        )
        assert pdf.status_code == 200, pdf.text
        assert pdf.content.startswith(b"%PDF")
        pdf_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(pdf.content)).pages
        )
        assert "ОБЩИЕ И ЦЕНТРАЛЬНЫЕ" in pdf_text.upper()
        assert "РАЗРЕШЕНО РУКОВОДИТЕЛЕМ ДЛЯ ТЕНДЕРА" not in pdf_text.upper()
        fixed = client.get(
            f"/api/projects/{project['id']}/quote", headers=admin_headers
        ).json()
        assert fixed["status"] == "fixed"
        assert fixed["revision"] == 1
        original_total = fixed["totals"]["grand_total"]

        appendix_preview = client.get(
            f"/api/projects/{project['id']}/documents/contract_appendix/preview",
            params={"token": token},
        )
        appendix_pdf = client.get(
            f"/api/projects/{project['id']}/documents/contract_appendix/pdf",
            headers=admin_headers,
        )
        appendix_docx = client.get(
            f"/api/projects/{project['id']}/documents/contract_appendix/docx",
            headers=admin_headers,
        )
        assert appendix_preview.status_code == 200
        assert appendix_pdf.status_code == 200
        assert appendix_docx.status_code == 200
        assert "Приложение к договору" in appendix_preview.text
        assert "2000 мм" in appendix_preview.text
        assert "2400 мм" in appendix_preview.text
        appendix_pdf_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(io.BytesIO(appendix_pdf.content)).pages
        )
        assert "ПРИЛОЖЕНИЕ К ДОГОВОРУ" in appendix_pdf_text.upper()
        with zipfile.ZipFile(io.BytesIO(appendix_docx.content)) as archive:
            appendix_word_xml = archive.read("word/document.xml").decode("utf-8")
        assert "ПРИЛОЖЕНИЕ К ДОГОВОРУ" in appendix_word_xml
        assert "2000 мм" in appendix_word_xml
        assert "2400 мм" in appendix_word_xml
        assert (
            client.get(
                f"/api/projects/{project['id']}/quote", headers=admin_headers
            ).json()["totals"]["grand_total"]
            == original_total
        )

        unchanged_refresh = client.post(
            f"/api/projects/{project['id']}/quote/refresh",
            headers=admin_headers,
        )
        assert unchanged_refresh.status_code == 200, unchanged_refresh.text
        assert unchanged_refresh.json()["revision"] == 1
        assert unchanged_refresh.json()["status"] == "fixed"
        assert unchanged_refresh.json()["totals"]["grand_total"] == original_total
        fixed_details = next(
            row["section_details"]
            for row in fixed["lines"]
            if row.get("section_details")
        )
        assert fixed_details["width_mm"] == "2000"

        db = SessionLocal()
        try:
            changed = db.get(
                models.CatalogPriceVersion,
                seed["version_by_sku"][seed["margin_sku"]],
            )
            replacement = models.CatalogPriceVersion(
                catalog_item_id=changed.catalog_item_id,
                cost=Decimal("999.00"),
                profile_markup_percent=changed.profile_markup_percent,
                profile_discount_percent=changed.profile_discount_percent,
                waste_markup_percent=changed.waste_markup_percent,
                construction_markup_percent=changed.construction_markup_percent,
                construction_discount_percent=changed.construction_discount_percent,
                category=changed.category,
                unit=changed.unit,
                min_margin_percent=changed.min_margin_percent,
                effective_from=datetime.utcnow(),
                created_at=datetime.utcnow(),
                created_by=admin_id,
                reason="Новая цена после фиксации",
            )
            db.add(replacement)
            changed_section = db.get(models.Section, section["id"])
            changed_section.width = 2100
            db.get(models.Project, project["id"]).updated_at = datetime.utcnow()
            db.commit()
            db.refresh(replacement)
            seed["versions"].append(replacement.id)
        finally:
            db.close()

        stale = client.get(
            f"/api/projects/{project['id']}/quote", headers=admin_headers
        ).json()
        assert stale["stale"] is True
        assert stale["totals"]["grand_total"] == original_total

        stale_word = client.get(
            f"/api/projects/{project['id']}/documents/commercial/docx",
            headers=admin_headers,
        )
        assert stale_word.status_code == 200
        with zipfile.ZipFile(io.BytesIO(stale_word.content)) as archive:
            stale_word_xml = archive.read("word/document.xml").decode("utf-8")
        assert "Расчёт устарел" in stale_word_xml
        assert "2000 мм" in stale_word_xml
        assert "2100 мм" not in stale_word_xml

        blocked_refresh = client.post(
            f"/api/projects/{project['id']}/quote/refresh",
            headers=admin_headers,
        )
        assert blocked_refresh.status_code == 409
        invalid_internal = client.get(
            f"/api/pricing/projects/{project['id']}", headers=admin_headers
        ).json()
        invalid_approval = invalid_internal["margin_approval"]
        assert invalid_approval["required"] is True
        assert invalid_approval["valid"] is False
        assert invalid_approval["target_revision"] == 2
        assert invalid_approval["approved_revision"] == 1
        assert invalid_internal["config"]["margin_override_comment"] == ""

        reapproved = client.put(
            f"/api/projects/{project['id']}/quote/overrides",
            headers=admin_headers,
            json={
                "overrides": [
                    {
                        "sku": seed["skipped_sku"],
                        "cost": "12",
                        "comment": "Разовая согласованная цена",
                    }
                ],
                "margin_override_comment": "Повторно согласовано для редакции 2",
            },
        )
        assert reapproved.status_code == 200, reapproved.text
        reapproved_internal = client.get(
            f"/api/pricing/projects/{project['id']}", headers=admin_headers
        ).json()
        assert reapproved_internal["margin_approval"]["valid"] is True
        assert reapproved_internal["margin_approval"]["target_revision"] == 2

        refreshed = client.post(
            f"/api/projects/{project['id']}/quote/refresh",
            headers=admin_headers,
        )
        assert refreshed.status_code == 200, refreshed.text
        assert refreshed.json()["revision"] == 2
        assert refreshed.json()["status"] == "fixed"
        assert refreshed.json()["totals"]["grand_total"] != original_total
        refreshed_details = next(
            row["section_details"]
            for row in refreshed.json()["lines"]
            if row.get("section_details")
        )
        assert refreshed_details["width_mm"] == "2100"
        refreshed_word = client.get(
            f"/api/projects/{project['id']}/documents/commercial/docx",
            headers=admin_headers,
        )
        assert refreshed_word.status_code == 200
        with zipfile.ZipFile(io.BytesIO(refreshed_word.content)) as archive:
            refreshed_word_xml = archive.read("word/document.xml").decode("utf-8")
        assert "2100 мм" in refreshed_word_xml
        assert "2400 мм" in refreshed_word_xml
    finally:
        _cleanup_quote_prices(seed)


def test_margin_approval_is_bound_to_every_pricing_input(client, admin_headers):
    admin_id = client.get("/api/auth/me", headers=admin_headers).json()["id"]
    dealer, dealer_password = _create_user(
        client,
        admin_headers,
        role="dealer",
    )
    dealer_headers = _login_headers(client, dealer["username"], dealer_password)
    project_id = None
    seed = None
    override_cost = "12"
    expected_revision = 1

    try:
        created = client.post(
            "/api/projects",
            headers=dealer_headers,
            json={"number": _unique("MARGIN-CONTEXT"), "customer": "Дилер"},
        )
        assert created.status_code == 201, created.text
        project_id = created.json()["id"]
        created_section = client.post(
            f"/api/projects/{project_id}/sections",
            headers=dealer_headers,
            json={
                "name": "Секция 1",
                "system": "СЛАЙД",
                "width": 2000,
                "height": 2400,
                "panels": 3,
                "quantity": 1,
                "rails": 3,
                "first_panel_inside": "Справа",
            },
        )
        assert created_section.status_code == 201, created_section.text
        section_id = created_section.json()["id"]
        seed = _seed_quote_prices(project_id, admin_id)

        def override_payload(comment: str | None = None):
            payload = {
                "overrides": [
                    {
                        "sku": seed["skipped_sku"],
                        "cost": override_cost,
                        "comment": "Разовая согласованная цена",
                    }
                ]
            }
            if comment is not None:
                payload["margin_override_comment"] = comment
            return payload

        approved = client.put(
            f"/api/projects/{project_id}/quote/overrides",
            headers=admin_headers,
            json=override_payload("Первичное согласование"),
        )
        assert approved.status_code == 200, approved.text

        db = SessionLocal()
        try:
            fixed = freeze_quote(
                db,
                db.get(models.Project, project_id),
                db.get(models.User, admin_id),
            )
            db.commit()
            assert fixed["revision"] == 1
            assert fixed["status"] == "fixed"
        finally:
            db.close()

        def require_reapproval(reason: str):
            nonlocal expected_revision
            internal = client.get(
                f"/api/pricing/projects/{project_id}", headers=admin_headers
            )
            assert internal.status_code == 200, internal.text
            approval = internal.json()["margin_approval"]
            assert approval["required"] is True, reason
            assert approval["valid"] is False, reason
            assert approval["target_revision"] == expected_revision + 1, reason
            assert internal.json()["config"]["margin_override_comment"] == ""

            blocked = client.post(
                f"/api/projects/{project_id}/quote/refresh",
                headers=admin_headers,
            )
            assert blocked.status_code == 409, (reason, blocked.text)
            still_fixed = client.get(
                f"/api/projects/{project_id}/quote", headers=admin_headers
            ).json()
            assert still_fixed["revision"] == expected_revision
            assert still_fixed["status"] == "fixed"

            reapproved = client.put(
                f"/api/projects/{project_id}/quote/overrides",
                headers=admin_headers,
                json=override_payload(f"Повторное согласование: {reason}"),
            )
            assert reapproved.status_code == 200, (reason, reapproved.text)
            approved_state = client.get(
                f"/api/pricing/projects/{project_id}", headers=admin_headers
            ).json()["margin_approval"]
            assert approved_state["valid"] is True, reason
            assert approved_state["target_revision"] == expected_revision + 1

            refreshed = client.post(
                f"/api/projects/{project_id}/quote/refresh",
                headers=admin_headers,
            )
            assert refreshed.status_code == 200, (reason, refreshed.text)
            expected_revision += 1
            assert refreshed.json()["revision"] == expected_revision
            assert refreshed.json()["status"] == "fixed"

        db = SessionLocal()
        try:
            current = db.get(
                models.CatalogPriceVersion,
                seed["version_by_sku"][seed["margin_sku"]],
            )
            replacement = models.CatalogPriceVersion(
                catalog_item_id=current.catalog_item_id,
                cost=Decimal(current.cost) + Decimal("1"),
                profile_markup_percent=current.profile_markup_percent,
                profile_discount_percent=current.profile_discount_percent,
                waste_markup_percent=current.waste_markup_percent,
                construction_markup_percent=current.construction_markup_percent,
                construction_discount_percent=current.construction_discount_percent,
                category=current.category,
                unit=current.unit,
                min_margin_percent=current.min_margin_percent,
                effective_from=datetime.utcnow(),
                created_at=datetime.utcnow(),
                created_by=admin_id,
                reason="Изменение цены для проверки контекста",
            )
            db.add(replacement)
            db.commit()
            db.refresh(replacement)
            seed["versions"].append(replacement.id)
        finally:
            db.close()
        require_reapproval("цена")

        dealer_terms = client.put(
            f"/api/pricing/dealers/{dealer['id']}",
            headers=admin_headers,
            json={
                "dealer_markup_percent": "5",
                "profile_discount_percent": "6",
                "construction_discount_percent": "1",
                "component_discount_percent": "7",
                "service_discount_percent": "2",
            },
        )
        assert dealer_terms.status_code == 200, dealer_terms.text
        require_reapproval("дилерские условия")

        db = SessionLocal()
        try:
            db.get(models.Section, section_id).width = 2100
            db.get(models.Project, project_id).updated_at = datetime.utcnow()
            db.commit()
        finally:
            db.close()
        require_reapproval("проект")

        service_change = client.put(
            f"/api/projects/{project_id}/quote/config",
            headers=admin_headers,
            json={
                "validity_days": 14,
                "manufacturing_term": "",
                "payment_terms": "",
                "services": [
                    {
                        "id": "delivery",
                        "name": "Доставка",
                        "quantity": "1",
                        "unit": "шт.",
                        "base_cost": "100",
                    }
                ],
            },
        )
        assert service_change.status_code == 200, service_change.text
        require_reapproval("услуга")

        override_cost = "13"
        override_change = client.put(
            f"/api/projects/{project_id}/quote/overrides",
            headers=admin_headers,
            json=override_payload(),
        )
        assert override_change.status_code == 200, override_change.text
        require_reapproval("разовая цена")
    finally:
        if project_id is not None:
            client.delete(f"/api/projects/{project_id}", headers=dealer_headers)
        if seed is not None:
            _cleanup_quote_prices(seed)
        db = SessionLocal()
        try:
            db.query(models.DealerPricingTerms).filter_by(user_id=dealer["id"]).delete()
            db.commit()
        finally:
            db.close()
        client.delete(f"/api/users/{dealer['id']}", headers=admin_headers)
