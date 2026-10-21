"""October 20: schematic side-profile scale and document-only RS3020 covers."""

from copy import deepcopy
import io
import re
from types import SimpleNamespace
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from lxml import html as lxml_html
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from engine.office_diagrams import section_diagrams
from engine.office_docx import build_project_docx, build_section_docx
from engine.office_xlsx import build_project_xlsx, build_section_xlsx
from engine.pdf import generate_pdf, render_pdf_html
from engine.project_documents import (
    CalculatedSection,
    _build_paint_pages,
    build_project_document_context,
    render_project_document_html,
)
from engine.slide_calc import calculate_slide
from schemas import SectionCreate


COVERS = "накладки замка RS3020"
LOCK = "ЗАМОК двухсторонний с ключом RS3020"


def section(**overrides):
    return SimpleNamespace(**SectionCreate(**{
        "name": "Секция 1", "system": "СЛАЙД",
        "width": 900, "height": 2400, "panels": 2, "quantity": 1,
        "rails": 3, "slide_rows": 1,
        "painting_type": "RAL стандарт", "ral_color": "7021 матовый",
        "profile_left_wall": True, "profile_right_wall": True,
        "profile_left_p_bar": True, "profile_right_p_bar": True,
        "profile_left_handle_bar": True, "profile_right_handle_bar": True,
        **overrides,
    }).model_dump())


def room_svg(s, calc):
    html = render_pdf_html(SimpleNamespace(number="B26-7-4303"), s, calc)
    return ET.fromstring(re.search(
        r'<svg data-office-diagram="slide-room".*?</svg>', html, re.S
    ).group())


@pytest.mark.parametrize("width,height", [(900, 2400), (4000, 2000), (10000, 2400)])
@pytest.mark.parametrize("rows,rails", [(1, 3), (1, 5), (2, 3), (2, 5)])
@pytest.mark.parametrize("mirror", [False, True])
def test_room_profiles_keep_physical_thickness_in_wide_frame(width, height, rows, rails, mirror):
    s = section(width=width, height=height, slide_rows=rows,
                panels=2 if rows == 1 else 4, rails=rails,
                first_panel_inside="Справа" if mirror else "Слева")
    # Different side assemblies also check their inward offsets and mirror anchors.
    short_side = "left" if mirror else "right"
    setattr(s, f"profile_{short_side}_p_bar", False)
    setattr(s, f"profile_{short_side}_handle_bar", False)
    calc = calculate_slide(s)
    before = deepcopy(calc)
    svg = room_svg(s, calc)
    assert svg.attrib["viewBox"] == "0 0 300 80"
    frame = next(node for node in svg if node.tag.endswith("rect"))
    assert [float(frame.get(key)) for key in ("x", "y", "width", "height")] == [10, 2, 280, 76]
    # Profiles fit the physical height and cannot grow beyond the width scale.
    scale = min(280 / width, 76 / height)
    for side in ("left", "right"):
        expected = [16] if side == short_side else [16, 18, 14, 13]
        parts = [node for node in svg if node.get("data-side") == side]
        assert [float(node.get("width")) for node in parts] == pytest.approx(
            [mm * scale for mm in expected]
        )
        offset = 0
        for node, mm in zip(parts, expected, strict=True):
            x = 10 + offset * scale if side == "left" else 290 - (offset + mm) * scale
            assert float(node.get("x")) == pytest.approx(x)
            offset += mm
    glass = [node for node in svg if node.get("data-glass-panel")]
    left = 10 + (16 if short_side == "left" else 61) * scale
    right = 290 - (16 if short_side == "right" else 61) * scale
    assert float(glass[0].get("x")) == pytest.approx(left)
    assert float(glass[-1].get("x")) + float(glass[-1].get("width")) == pytest.approx(right)
    assert calc == before


@pytest.mark.parametrize("flags,expected", [
    ({}, []),
    ({"wall": True}, [16]),
    ({"lock_bar": True, "handle_bar": True}, [50, 14, 13]),
    ({"wall": True, "p_bar": True}, [16, 25]),
    ({"bubble": True}, [10]),
])
def test_room_optional_side_assemblies(flags, expected):
    overrides = {
        f"profile_{side}_{flag}": flags.get(flag, False)
        for side in ("left", "right")
        for flag in ("wall", "lock_bar", "p_bar", "handle_bar", "bubble")
    }
    s = section(**overrides)
    svg = room_svg(s, calculate_slide(s))
    for side in ("left", "right"):
        parts = [node for node in svg if node.get("data-side") == side]
        assert [float(node.get("width")) for node in parts] == pytest.approx(
            [mm * 76 / 2400 for mm in expected]
        )


def test_screenshot_case_keeps_glass_labels_and_office_uses_same_svg():
    s = section()
    calc = calculate_slide(s)
    assert [(g.width_mm, g.height_mm, g.qty) for g in calc.glass] == [(411, 2294, 2)]
    svg = room_svg(s, calc)
    assert list(svg.itertext()).count("411 (427)") == 2
    parts = [node for node in svg if node.get("data-side") == "left"]
    assert sum(float(node.get("width")) for node in parts) == pytest.approx(61 * 76 / 2400)
    diagrams = section_diagrams(s, calc)
    assert len(diagrams) == 2
    for export in (build_section_docx, build_section_xlsx):
        with ZipFile(io.BytesIO(export(SimpleNamespace(number="QA"), s, calc))) as archive:
            images = [archive.read(name) for name in archive.namelist() if "/media/" in name]
            assert diagrams[0][1] in images
    assert COVERS not in render_pdf_html(SimpleNamespace(number="QA"), s, calc)


@pytest.mark.parametrize("rows", [1, 2])
@pytest.mark.parametrize("locks,quantity", [(0, 1), (1, 1), (1, 3), (2, 1), (2, 3)])
def test_covers_are_twice_calculated_locks_without_double_section_multiplier(rows, locks, quantity):
    s = section(slide_rows=rows, panels=2 if rows == 1 else 4,
                quantity=quantity, lock_left=LOCK if locks else "Без",
                lock_right=LOCK if locks == 2 else "Без")
    project = SimpleNamespace(number="QA")
    expected = locks * quantity * 2
    paint = build_project_document_context(project, [s], "paint")
    delivery = build_project_document_context(project, [s], "delivery")
    for rows in ([row for page in paint["paint_pages"] for row in page["rows"]],
                 delivery["delivery_item2_rows"]):
        found = [row for row in rows if row["name"] == COVERS]
        assert len(found) == (1 if expected else 0)
        if found:
            assert found[0]["article"] == ""
            assert found[0]["qty"] == expected


def test_covers_group_by_color_and_do_not_change_profile_lengths_or_calc():
    s = section(lock_left=LOCK, lock_right=LOCK, quantity=3)
    calc = calculate_slide(s)
    before = deepcopy(calc)
    pages = _build_paint_pages([CalculatedSection(order=1, section=s, calc=calc)])
    assert calc == before
    without_locks = deepcopy(calc)
    without_locks.hardware = [item for item in calc.hardware if item.article != "RS3020"]
    reference = _build_paint_pages([CalculatedSection(order=1, section=s, calc=without_locks)])
    assert pages[0]["total_m"] == reference[0]["total_m"]
    assert pages[0]["total_qty"] == reference[0]["total_qty"] + 12
    assert [row for row in pages[0]["rows"] if row["name"] != COVERS] == reference[0]["rows"]

    sections = [s, section(lock_left=LOCK),
                section(lock_left=LOCK, ral_color="9016")]
    project = SimpleNamespace(number="QA")
    paint = build_project_document_context(project, sections, "paint")
    assert {
        page["color"]: next(row["qty"] for row in page["rows"] if row["name"] == COVERS)
        for page in paint["paint_pages"]
    } == {"RAL 7021 матовый": 14, "RAL 9016": 2}
    delivery = build_project_document_context(project, sections, "delivery")
    assert {row["color"]: row["qty"] for row in delivery["delivery_item2_rows"]
            if row["name"] == COVERS} == {"RAL 7021 матовый": 14, "RAL 9016": 2}


@pytest.mark.parametrize("installed", ["installed", "not_installed"])
@pytest.mark.parametrize("stages,stage,present", [(1, 1, True), (2, 1, False), (2, 2, True)])
def test_covers_delivery_follows_hardware_stage_even_when_installed(installed, stages, stage, present):
    project = SimpleNamespace(number="QA", production_stages=stages,
                              current_stage=stage, hardware_installation=installed)
    context = build_project_document_context(project, [section(lock_left=LOCK)], "delivery")
    found = [row for row in context["delivery_item2_rows"] if row["name"] == COVERS]
    assert bool(found) == present
    if present:
        assert found[0]["qty_display"] == "2 шт"
        assert found[0]["color"] == "RAL 7021 матовый"


def test_covers_require_exact_lock_article_and_do_not_invent_ral():
    project = SimpleNamespace(number="QA")
    for s in (section(lock_left="ЗАМОК-ЗАЩЕЛКА 1стор RS3018"),
              section(handle_left="Ручка-скоба 600мм RS30201")):
        assert COVERS not in render_project_document_html(project, [s], "paint")
        assert COVERS not in render_project_document_html(project, [s], "delivery")
    s = section(lock_left=LOCK, painting_type="", ral_color="")
    for doc_type in ("paint", "delivery"):
        html = render_project_document_html(project, [s], doc_type)
        assert COVERS in html
        assert "RAL 7021" not in html


def test_covers_paint_exports_keep_blank_article_and_dimensions_with_image():
    s = section(lock_left=LOCK, lock_right=LOCK, quantity=3)
    project = SimpleNamespace(number="QA")
    html = lxml_html.fromstring(render_project_document_html(project, [s], "paint"))
    row = html.xpath(f'//tr[td//div[text()="{COVERS}"]]')[0]
    cells = row.findall("td")
    assert [cell.text_content().strip() for cell in cells][2:] == ["12", "", "", ""]
    assert cells[0].text_content().strip() == ""
    assert cells[1].xpath('.//img[starts-with(@src,"data:image/png;base64,")]')

    doc = Document(io.BytesIO(build_project_docx(project, [s], "paint")))
    row = next(row for table in doc.tables for row in table.rows if COVERS in row.cells[1].text)
    assert [cell.text.strip() for cell in row.cells] == ["", COVERS, "12", "", "", ""]
    assert row.cells[1]._tc.xpath(".//w:drawing")

    for cached in (False, True):
        workbook = load_workbook(io.BytesIO(build_project_xlsx(project, [s], "paint")), data_only=cached)
        sheet = workbook.active
        row = next(row for row in sheet if row[1].value == COVERS)
        assert [cell.value for cell in row[:6]] == [None, COVERS, 12, None, None, None]
        assert sheet._images
        total = next(row for row in sheet if row[0].value == "Итого")
        if cached:
            context = build_project_document_context(project, [s], "paint")["paint_pages"][0]
            assert total[2].value == context["total_qty"]
            assert total[5].value == context["total_m"]
        else:
            assert total[2].value.startswith("=SUM(")
            assert total[5].data_type == "f" and "SUM(" in total[5].value

    workbook = load_workbook(io.BytesIO(build_project_xlsx(project, [s], "delivery")))
    row = next(row for row in workbook.active if any(COVERS in str(cell.value) for cell in row))
    values = [cell.value for cell in row]
    assert 12 in values
    assert any("RAL 7021 матовый" in str(value) for value in values)


def test_covers_do_not_enter_production_sketch_hardware_order_or_calculation():
    s = section(lock_left=LOCK)
    project = SimpleNamespace(number="QA")
    calc = calculate_slide(s)
    assert not any(item.name == COVERS for item in calc.hardware)
    assert COVERS not in render_pdf_html(project, s, calc)
    for doc_type in ("sketch", "hardware_order"):
        assert COVERS not in render_project_document_html(project, [s], doc_type)


@pytest.mark.parametrize("doc_type", ["paint", "delivery"])
def test_covers_pdf_contains_name_color_and_quantity(doc_type):
    s = section(lock_left=LOCK, lock_right=LOCK, quantity=3)
    html = render_project_document_html(SimpleNamespace(number="QA"), [s], doc_type, is_pdf=True)
    pdf = PdfReader(io.BytesIO(generate_pdf(html)))
    text = "\n".join(page.extract_text(extraction_mode="layout") for page in pdf.pages)
    assert COVERS in " ".join(text.split())
    assert "7021" in text
    assert re.search(r"\b12\b", text)
