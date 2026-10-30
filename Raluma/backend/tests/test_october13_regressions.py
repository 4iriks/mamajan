"""Only the delivery amendments of October 13; other documents stay unchanged."""

from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from engine.cs_calc import calculate_cs
from engine.office_xlsx import build_project_xlsx
from engine.project_documents import (
    build_project_document_context,
    render_project_document_html,
)
from schemas import SectionCreate

# Exports resolve catalogue entries; initialize/clean the test catalogue through
# the normal application lifespan before rendering, regardless of test order.
pytestmark = pytest.mark.usefixtures("client")


@pytest.mark.parametrize(
    "stages,stage,pads_present,tape_present",
    [
        (1, 1, True, True),
        (2, 1, True, False),
        (2, 2, False, True),
    ],
)
@pytest.mark.parametrize("quantity,divisions", [(1, 0), (3, 2)])
@pytest.mark.parametrize("custom_pad", [False, True])
def test_cs_pads_ship_with_clamp_once(
    stages, stage, pads_present, tape_present, quantity, divisions, custom_pad
):
    section = SimpleNamespace(
        **SectionCreate(
            name="Секция 1",
            system="ЦС",
            width=2240,
            height=600,
            quantity=quantity,
            cs_config={
                "version": 2,
                "vertices": [
                    {"x": 0, "y": 0},
                    {"x": 2240, "y": 0},
                    {"x": 2240, "y": 600},
                    {"x": 0, "y": 600},
                ],
                "vertical": {"count": divisions, "mode": "equal"},
            },
        ).model_dump()
    )
    sku = "CUSTOM-PAD" if custom_pad else "CS-PVC-PAD"
    if custom_pad:
        section.cs_system = SimpleNamespace(
            glass_pad=SimpleNamespace(sku=sku, name="Подкладка заказная")
        )
    project = SimpleNamespace(
        number="QA", production_stages=stages, current_stage=stage
    )
    before = deepcopy(section.__dict__)
    calculated = calculate_cs(section, system=getattr(section, "cs_system", None))
    context = build_project_document_context(project, [section], "delivery")
    hardware = context["delivery_item2_rows"]
    pads = [row for row in hardware if row["article"] == sku]
    clamps = [row for row in hardware if row["article"] == "Т40Т"]
    tape = [row for row in hardware if "скотч" in row["name"].lower()]
    assert bool(pads) == bool(clamps) == pads_present
    assert bool(tape) == (tape_present and divisions > 0)
    if pads_present:
        assert len(pads) == 1
        assert pads[0]["qty"] == len(calculated["panes"]) * 2 * quantity
    html = render_project_document_html(project, [section], "delivery", is_pdf=True)
    assert (sku in html) == pads_present
    sheet = load_workbook(
        BytesIO(build_project_xlsx(project, [section], "delivery"))
    ).active
    xlsx_pads = [row for row in sheet.values if str(row[1]).startswith(sku + " — ")]
    assert bool(xlsx_pads) == pads_present
    if pads_present:
        assert xlsx_pads[0][4] == pads[0]["qty"]
    assert section.__dict__ == before


def test_delivery_lengths_have_names_without_images_and_keep_place_keys():
    base = dict(
        sku="BOX-40",
        name="Бокс 40",
        color="9005",
        finishName="RAL стандарт",
        unit="м",
        qty=2,
        image="T40K-section.jpg",
        deliveryStage="both",
    )
    extras = [{**base, "size": length} for length in ("1000", "1500", "2000")]
    extras.append({**base, "size": "1000", "color": "9016"})
    project = SimpleNamespace(number="QA", extra_components=extras)
    before = deepcopy(extras)
    context = build_project_document_context(project, [], "delivery")
    rows = context["delivery_project_extra_rows"]
    assert len(rows) == 4
    assert len({row["place_key"] for row in rows}) == 4
    for is_pdf in (False, True):
        html = render_project_document_html(project, [], "delivery", is_pdf=is_pdf)
        assert html.count("BOX-40") == html.count("Бокс 40") == 2
        assert html.count('<td class="extra-size">') == 4
        assert "<img" not in html
        if not is_pdf:
            assert all(row["place_key"] in html for row in rows)
    data = build_project_xlsx(project, [], "delivery")
    sheet = load_workbook(BytesIO(data)).active
    headings = [row[1] for row in sheet.values if "BOX-40" in str(row[1])]
    assert len(headings) == 2
    assert any("9005" in heading for heading in headings)
    assert any("9016" in heading for heading in headings)
    xlsx_rows = [row for row in sheet.values if str(row[3]).endswith(" мм")]
    assert len(xlsx_rows) == 4
    assert all(row[4] == 2 for row in xlsx_rows)
    with ZipFile(BytesIO(data)) as archive:
        assert not any(path.startswith("xl/media/") for path in archive.namelist())
    assert extras == before
