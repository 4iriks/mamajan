from collections import Counter
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

from docx import Document
from pypdf import PdfReader
import pytest

from engine.pdf import generate_pdf
from engine.project_documents import (
    build_project_document_context,
    render_project_document_html,
)
from engine.office_docx import _build_sketch_docx
from engine.sketch_layout import fit_sketch_context
from schemas import SectionCreate


def section(**overrides):
    return SimpleNamespace(**SectionCreate(**overrides).model_dump())


def dense_slide(panels=10):
    return section(
        system="СЛАЙД",
        name="Секция 2",
        order=2,
        width=7704,
        height=2633,
        slide_rows=2,
        rails=5,
        panels=panels,
        threshold="Стандартный окраш",
        painting_type="RAL стандарт",
        ral_color="8019 МАТОВЫЙ",
        profile_left_wall=True,
        profile_right_wall=True,
        profile_left_p_bar=True,
        profile_right_p_bar=True,
        profile_left_bubble=True,
        profile_right_bubble=True,
        center_handle="Ручки-профиль RS112 (2шт)",
        center_lock="Накладная защёлка RS206",
        comments=("Проверить расположение профилей и направление открывания. " * 18)
        + "КОНЕЦ ПРИМЕЧАНИЯ",
    )


@pytest.mark.parametrize("panels", [4, 6, 8, 10])
def test_dense_product_stays_on_one_pdf_page_with_all_glass_rows(panels):
    project = SimpleNamespace(number="FIT-CHECK")
    item = dense_slide(panels)
    before = deepcopy(vars(item))
    context = build_project_document_context(project, [item], "sketch")
    html = render_project_document_html(project, [item], "sketch", is_pdf=True)
    pdf = PdfReader(BytesIO(generate_pdf(html)))
    assert len(pdf.pages) == 1
    text = " ".join(pdf.pages[0].extract_text().split())
    assert "КОНЕЦ ПРИМЕЧАНИЯ" in text and "РАЗМЕРЫ СТЕКОЛ" in text.upper()
    expected = Counter(p.position for p in context["sections"][0]["panels"])
    for position, count in expected.items():
        assert text.count(position) == count
    assert vars(item) == before


def test_each_product_has_its_own_layout_and_word_keeps_last_row_separate():
    project = SimpleNamespace(number="MIXED-FIT")
    items = [
        section(system="ЦС", name="Секция 1", order=1, width=1000, height=600),
        dense_slide(),
    ]
    context = build_project_document_context(project, items, "sketch")
    layouts = fit_sketch_context(context)
    assert layouts[0].scale == 1 and not layouts[0].compact
    assert layouts[1].compact and layouts[1].scale < 1
    html = render_project_document_html(project, items, "sketch", is_pdf=True)
    pdf = PdfReader(BytesIO(generate_pdf(html)))
    assert len(pdf.pages) == 2
    for index, page in enumerate(pdf.pages, 1):
        assert f"Изделие № {index}" in page.extract_text()
        assert "РАЗМЕРЫ СТЕКОЛ" in page.extract_text().upper()
    word = Document(BytesIO(_build_sketch_docx(context)))
    last = word.tables[-1]
    assert len(last.rows) == 11
    # The grid must match cell widths; otherwise Word wraps filling names
    # into narrow equal-width columns and spills the glass table.
    assert [round(column.width.mm) for column in last.columns] == [
        12,
        32,
        70,
        25,
        25,
        18,
    ]
    assert all(
        p.paragraph_format.keep_with_next is False
        for c in last.rows[-1].cells
        for p in c.paragraphs
    )
    assert "КОНЕЦ ПРИМЕЧАНИЯ" in " ".join(p.text for p in word.paragraphs)


@pytest.mark.parametrize("system", ["КНИЖКА", "ЛИФТ"])
def test_other_systems_keep_glass_on_the_product_page(system):
    project = SimpleNamespace(number="SYSTEM-FIT")
    item = section(name="Секция 1", system=system, width=4000, height=2700, panels=4)
    html = render_project_document_html(project, [item], "sketch", is_pdf=True)
    pdf = PdfReader(BytesIO(generate_pdf(html)))
    assert len(pdf.pages) == 1
    assert "РАЗМЕРЫ СТЕКОЛ" in pdf.pages[0].extract_text().upper()
