"""October 21: real stock layouts shared by quotes and fulfilment documents."""

from collections import Counter
from decimal import Decimal
import io
import json
from types import SimpleNamespace
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
import pytest

import models
import schemas
from engine.cs_calc import calculate_cs
from engine.cs_cover_cutting import build_cover_plan, cover_cuts, section_allocation_m
from engine.project_documents import (
    build_project_document_context,
    render_project_document_html,
)
from engine.office_docx import build_project_docx
from engine.office_xlsx import build_project_xlsx
from engine.quote_pricing import (
    calculate_quote,
    get_or_create_quote_state,
    public_quote,
    quote_is_stale,
    refresh_quote_revision,
    safe_public_payload,
)
from engine.construction_pricing import migrate_construction_pricing
from tests.test_construction_pricing import setup, add_cs  # noqa: F401


def sections():
    return [
        SimpleNamespace(
            **schemas.SectionCreate(
                name=f"Секция {i}",
                order=i,
                system="ЦС",
                width=w,
                height=h,
                painting_type="RAL муар",
                ral_color="8025 МУАР",
            ).model_dump()
        )
        for i, w, h in [(4, 1865, 890), (5, 2125, 890), (6, 1520, 2385)]
    ]


@pytest.mark.parametrize(
    "length,expected",
    [
        (200, [250]),
        (2950, [3000]),
        (2951, [3000, 500]),
        (3010, [3000, 500]),
        (3400, [3000, 500]),
        (3401, [3000, 501]),
        (5900, [3000, 3000]),
        (5901, [3000, 3000, 500]),
    ],
)
def test_long_cover_tail_preserves_joint_and_minimum(length, expected):
    cuts = cover_cuts([length])
    for face in range(2):
        row = [cut for cut in cuts if cut["face"] == face]
        assert [cut["blank_mm"] for cut in row] == expected
        assert sum(cut["clean_mm"] for cut in row) == length
        assert all(cut["clean_mm"] > 0 for cut in row)


def test_control_individual_and_shared_cutting_and_allocations():
    entries = [(s, calculate_cs(s)) for s in sections()]
    assert [
        next(p["pieces"] for p in c["profiles"] if p["role"] == "cover")
        for s, c in entries
    ] == [4, 6, 8]
    plan = build_cover_plan(entries)
    (group,) = plan["groups"]
    assert group["stock_count"] == 16
    assert group["total_length_mm"] == 48000
    assert Counter(tuple(c["blank_mm"] for c in b) for b in group["stocks"]) == {
        (2435,): 4,
        (2175,): 4,
        (1915, 940): 4,
        (1570, 940): 4,
    }
    assert sum(section_allocation_m(plan, i) for i in range(3)) == 48
    assert section_allocation_m(plan, 0) == 12
    assert section_allocation_m(plan, 1) == Decimal("16.494024")
    assert build_cover_plan(entries) == plan


def test_quantities_finishes_edges_and_feasible_stock():
    rows = sections()
    rows[0].quantity = 2
    rows[1].ral_color = "9016"
    rows[2].painting_type = "Анод/неокрас"
    plan = build_cover_plan([(s, calculate_cs(s)) for s in rows])
    assert [g["stock_count"] for g in plan["groups"]] == [8, 6, 8]
    assert [g["requires_paint"] for g in plan["groups"]] == [True, True, False]
    for group in plan["groups"]:
        assert (
            sum(map(Decimal, group["allocations_m"].values())) * 1000
            == group["total_length_mm"]
        )
        assert all(sum(c["blank_mm"] for c in b) <= 3000 for b in group["stocks"])
    rows[0].cs_config = {"edgeTreatments": ["bubble", "none", "none", "none"]}
    assert not build_cover_plan([(rows[0], calculate_cs(rows[0]))])["groups"]


def test_repeated_small_sections_share_stock_and_different_profiles_do_not():
    row = sections()[0]
    row.width = row.height = 400
    row.quantity = 2
    calc = calculate_cs(row)
    assert next(p["pieces"] for p in calc["profiles"] if p["role"] == "cover") == 4
    plan = build_cover_plan([(row, calc)])
    assert plan["groups"][0]["stock_count"] == 3
    assert section_allocation_m(plan, 0) == 9
    assert Counter(
        c["copy_index"] for stock in plan["groups"][0]["stocks"] for c in stock
    ) == {0: 8, 1: 8}
    other = calculate_cs(row)
    next(p for p in other["profiles"] if p["role"] == "cover")["article"] = (
        "OTHER-COVER"
    )
    assert len(build_cover_plan([(row, calc), (row, other)])["groups"]) == 2


def test_documents_share_order_count_and_delivery_stage():
    project = SimpleNamespace(number="В26-5-4306", production_stages=1, current_stage=1)
    rows = sections()
    paint = build_project_document_context(project, rows, "paint")
    painted = [
        r for p in paint["paint_pages"] for r in p["rows"] if r["article"] == "Т40К"
    ]
    assert [(r["qty"], r["clean"], r["allowance"], r["total_m"]) for r in painted] == [
        (16, 3000, 3000, 48)
    ]
    delivery = build_project_document_context(project, rows, "delivery")
    assert [
        r["qty"] for r in delivery["delivery_item2_rows"] if r["article"] == "Т40К"
    ] == [16]
    assert "16" in render_project_document_html(project, rows, "hardware_order")
    doc = Document(io.BytesIO(build_project_docx(project, rows, "paint")))
    cover = next(r for t in doc.tables for r in t.rows if r.cells[0].text == "Т40К")
    assert cover.cells[2].text == "16"
    for kind in ("paint", "delivery", "hardware_order"):
        book = load_workbook(
            io.BytesIO(build_project_xlsx(project, rows, kind)), data_only=True
        )
        found = [list(r) for sheet in book for r in sheet.values if "Т40К" in r]
        assert found and any(16 in r for r in found), kind
    project.production_stages = 2
    project.current_stage = 2
    delivery = build_project_document_context(project, rows, "delivery")
    assert not [r for r in delivery["delivery_item2_rows"] if r["article"] == "Т40К"]


def test_quote_prices_shared_stock_and_preserves_fixed_snapshots(request):
    db, actor, _ = request.getfixturevalue("setup")
    first, second, third = sections()
    project, section, kit = add_cs(
        db,
        actor,
        width=first.width,
        height=first.height,
        painting_type=first.painting_type,
        ral_color=first.ral_color,
    )
    for source in (second, third):
        values = vars(source).copy()
        values["cs_system_id"] = kit.id
        db.add(models.Section(project_id=project.id, **values))
    db.flush()
    db.expire(project, ["sections"])
    migrate_construction_pricing(db)
    db.flush()
    state = get_or_create_quote_state(db, project)
    public, internal, signature = calculate_quote(db, project, state)
    assert public["export_allowed"], public["pricing_issues"]
    cover_rows = [
        r for s in internal["sections"] for r in s["bom"] if r.get("role") == "cover"
    ]
    assert sum(Decimal(r["quantity"]) for r in cover_rows) == 48
    assert public["cs_cover_cutting"]["groups"][0]["stock_count"] == 16
    assert safe_public_payload(public)["cs_cover_cutting"] == public["cs_cover_cutting"]
    for line in public["lines"]:
        cover = next(
            r
            for r in line["section_details"]["cs_geometry"]["profiles"]
            if r["role"] == "cover"
        )
        assert "pieces" not in cover and "cut_lengths_mm" not in cover
        assert "allocation_m" in cover
    state.source_signature = "legacy-calculation"
    draft = public_quote(db, project)
    assert not draft["stale"]
    assert draft["cs_cover_cutting"]["groups"][0]["stock_count"] == 16
    state.public_payload = json.dumps(public, ensure_ascii=False)
    state.internal_payload = json.dumps(internal, ensure_ascii=False)
    state.source_signature = signature
    state.status = "fixed"
    original = state.public_payload
    project.sections[0].width += 100
    db.flush()
    assert quote_is_stale(db, project, state)
    public_quote(db, project)
    assert state.public_payload == original
    updated = refresh_quote_revision(db, project, actor)
    assert updated["revision"] == 2 and not updated["stale"]
    archive = (
        db.query(models.ProjectQuoteRevision)
        .filter_by(project_id=project.id, revision=1)
        .one()
    )
    assert archive.public_payload == original


def test_cs_whole_glass_dimensions_tape_and_separate_delivery_group():
    rows = sections()
    rows[-1].cs_config = {"vertical": {"count": 1}}
    calc = calculate_cs(rows[-1])
    assert [(p["width_mm"], p["height_mm"]) for p in calc["panes"]] == [(732, 2333)] * 2
    assert calc["glass_area_m2"] == pytest.approx(732 * 2333 * 2 / 1e6)
    assert next(p["qty"] for p in calc["hardware"] if p["role"] == "tape") == 5
    assert calc["warnings"]  # Explanations remain available inside the program.
    slide = SimpleNamespace(
        **schemas.SectionCreate(
            name="Секция 1",
            system="СЛАЙД",
            width=1870,
            height=2500,
            order=1,
            painting_type=rows[0].painting_type,
            ral_color=rows[0].ral_color,
        ).model_dump()
    )
    project = SimpleNamespace(number="В26-5-4306")
    context = build_project_document_context(project, [slide, *rows], "delivery")
    groups = [row for row in context["delivery_item1_rows"] if row["kind"] == "glass"]
    cs = next(g for g in groups if g["system"] == "ЦС")
    assert cs["name"] == "СТЕКЛО(ЦС), РАЗМЕРЫ СТЕКОЛ:"
    assert cs["color"] == ""
    assert [r["marking"] for r in cs["rows"]] == ["4,1", "5,1", "6,1"]
    assert cs["rows"][-1]["width"] == 732
    assert any(g["name"].startswith("СТЕКЛЯННЫЕ ПАНЕЛИ, RAL") for g in groups)
    html = render_project_document_html(project, [slide, *rows], "delivery")
    assert "732.0" not in html and "732,5" not in html
    book = load_workbook(
        io.BytesIO(build_project_xlsx(project, [slide, *rows], "delivery")),
        data_only=True,
    )
    glass_rows = {
        r[2]: r
        for sheet in book
        for r in sheet.values
        if r[2] in {"1,1", "4,1", "5,1", "6,1"}
    }
    assert glass_rows["1,1"][1] == "Стеклянная панель"
    assert "8025" in glass_rows["1,1"][3]
    for marking, size in [
        ("4,1", "1813 × 838 мм"),
        ("5,1", "2073 × 838 мм"),
        ("6,1", "732 × 2333 мм"),
    ]:
        row = glass_rows[marking]
        assert row[1] == "СТЕКЛО(ЦС)"
        assert row[3] == f"10ММ ЗАКАЛЕННОЕ ПРОЗРАЧНОЕ\n{size}"


@pytest.mark.parametrize(
    "kind", ["glass", "paint", "delivery", "hardware_order", "sketch", "section"]
)
def test_system_warnings_do_not_leak_into_printed_documents(kind):
    from engine.cs_calc import cs_result_namespace
    from engine.office_docx import build_section_docx
    from engine.office_xlsx import build_section_xlsx
    from engine.pdf import render_pdf_html

    project = SimpleNamespace(number="В26-5-4306")
    row = sections()[-1]
    row.cs_config = {"vertical": {"count": 1}}
    row.comments = "Особая отметка заказчика"
    calc = cs_result_namespace(calculate_cs(row))
    texts = []
    archives = []
    if kind == "section":
        texts.append(render_pdf_html(project, row, calc))
        archives += [
            build_section_docx(project, row, calc),
            build_section_xlsx(project, row, calc),
        ]
    else:
        texts.append(render_project_document_html(project, [row], kind, is_pdf=True))
        if kind != "delivery":
            archives.append(build_project_docx(project, [row], kind))
        if kind != "sketch":
            archives.append(build_project_xlsx(project, [row], kind))
    for data in archives:
        with ZipFile(io.BytesIO(data)) as archive:
            texts.append(
                "\n".join(
                    archive.read(name).decode("utf-8")
                    for name in archive.namelist()
                    if name.endswith(".xml")
                )
            )
    for body in texts:
        for warning in [
            *calc.warnings,
            "предварительная ведомость",
            "Контур по схеме; размеры габаритные",
        ]:
            assert warning not in body
    if kind in {"sketch", "section"}:
        # Section XLSX never had a free-form comments field; HTML and Word do.
        assert all("Особая отметка заказчика" in body for body in texts[:2])
